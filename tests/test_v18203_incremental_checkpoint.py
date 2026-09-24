"""V18.20.3.0 characterization tests — incremental chapter checkpointing.

User requirement: "for example book is 10 chapter, 9 chapter is done, the
app crashes at the 10th chapter, we must have 9 chapters ready and
untouched by this crash, because they're already done and saved."

Covers:
- ``render_summary_markdown``: byte-parity with the end-of-run writer
  (year styling, additive sections, COMPLETE sentinel)
- ``ChapterCheckpointSaver``: immediate per-chapter save (md + docx +
  manifest), final-position numbering, write-once "untouched" guarantee,
  failed-chapter skip, best-effort exception containment
- crash-recovery integration: ``chapter_resume.parse_summary_file``
  accepts the checkpoint files (the V7.7.0 resume system finds them)
- ``Checkpoint.mark_progress``: records chapters_done while KEEPING the
  phase resumable (A15 invariant)
- ``Summarizer.summarize_chapters(chapter_done_callback=…)``: fires per
  chapter in BOTH sequential and parallel paths; callback exceptions
  never kill the loop
- settings default ``summarisation.incremental_chapter_save`` = True
"""

from __future__ import annotations

import json
import re
import sys
import types
from pathlib import Path
from unittest.mock import MagicMock

import pytest

# Stub the optional heavy deps exactly like tests/test_v18202_system_fixes.py
# does (markitdown is not installed in this sandbox; the stub keeps the
# eager core.pipeline→core.pdf import chain loadable).
if "markitdown" not in sys.modules:
    _md = types.ModuleType("markitdown")
    _md.MarkItDown = MagicMock()
    sys.modules["markitdown"] = _md

PROJECT_ROOT = Path(__file__).resolve().parent.parent


# ═══════════════════════════════════════════════════════════════════════════
# Test scaffolding — stub converter (no Qt / no LLM / no PDF deps)
# ═══════════════════════════════════════════════════════════════════════════


class _StubPaths:
    def __init__(self, root: Path):
        self.root = root
        self.summaries_dir = root / "summaries"
        self.book_summary = self.summaries_dir / "Book.summary.md"


class _StubCheckpoint:
    def __init__(self):
        self.calls = []

    def mark_progress(self, phase, **metadata):
        self.calls.append((phase, dict(metadata)))


class _StubRequest:
    generate_docx = True
    docx_author = "Elixir"
    docx_include_toc = True

    def __init__(self, p: Path):
        self.input_path = p


class _StubConverter:
    """Minimal converter surface used by ChapterCheckpointSaver."""

    def __init__(self, root: Path):
        self.output_dir = root
        self.paths = _StubPaths(root)
        self.checkpoint = _StubCheckpoint()
        self.request = _StubRequest(root / "book.pdf")
        self.settings = {}
        self.log_lines = []

    def _log(self, msg, level="INFO"):
        self.log_lines.append((level, msg))

    def _sanitize_filename(self, heading):
        safe = re.sub(r"[^\w\s\-]", "", heading or "")
        return re.sub(r"\s+", "_", safe.strip())[:80]


def _mk_summary(heading: str, n_bullets: int = 3, success: bool = True):
    from core.models import SummaryResult

    return SummaryResult(
        chapter_heading=heading,
        executive_summary=f"Executive summary for {heading} (1980) with a year.",
        thorough_summary=f"Thorough summary for {heading}.",
        bullet_points=[f"Bullet {i} about {heading} (1975)." for i in range(n_bullets)],
        success=success,
        error_msg=None if success else "boom",
    )


@pytest.fixture
def stub_env(tmp_path):
    root = tmp_path / "MyBook"
    root.mkdir()
    conv = _StubConverter(root)
    chapters = {
        "Chapter One": {"heading": "Chapter One", "content": "Content one. " * 60},
        "Chapter Two": {"heading": "Chapter Two", "content": "Content two. " * 60},
        "Chapter Three": {"heading": "Chapter Three", "content": "Content three. " * 60},
    }
    positions = {"Chapter One": 1, "Chapter Two": 2, "Chapter Three": 3}
    return root, conv, chapters, positions


# ═══════════════════════════════════════════════════════════════════════════
# 1) render_summary_markdown
# ═══════════════════════════════════════════════════════════════════════════


class TestRenderSummaryMarkdown:
    def test_format_matches_legacy_writer(self):
        """Byte-parity with the pre-V18.20.3.0 inline writer format."""
        from core.pipeline.chapter_checkpoint import render_summary_markdown

        summ = _mk_summary("Test Chapter", n_bullets=2)
        md = render_summary_markdown(summ)
        assert md.startswith("# Summary: Test Chapter\n\n## Executive Summary\n\n")
        assert "## Key Takeaways\n\n- Bullet 0 about Test Chapter" in md
        assert "## Thorough Summary\n\n" in md
        assert md.rstrip().endswith("<!-- COMPLETE -->")

    def test_year_styling_applied(self):
        from core.pipeline.chapter_checkpoint import render_summary_markdown

        md = render_summary_markdown(_mk_summary("Ch", 1))
        # V0.4.26 (audit backlog #5): format_years now emits markdown italic
        # (*YYYY*) instead of raw HTML <em> (which leaked into the markdown).
        assert "*1975*" in md
        assert "*1980*" in md

    def test_additive_sections_rendered_only_when_non_empty(self):
        from core.pipeline.chapter_checkpoint import render_summary_markdown

        s = _mk_summary("Ch", 1)
        s.key_distinctions = "KD text"
        s.theoretical_tensions = "TT text"
        s.scholar_reflections = ""
        md = render_summary_markdown(s)
        assert "## Key Distinctions\n\nKD text" in md
        assert "## Theoretical Tensions\n\nTT text" in md
        assert "Scholar" not in md

    def test_none_additive_sections_safe(self):
        from core.pipeline.chapter_checkpoint import render_summary_markdown

        md = render_summary_markdown(_mk_summary("Ch", 1))
        assert "Key Distinctions" not in md
        assert "Theoretical Tensions" not in md


# ═══════════════════════════════════════════════════════════════════════════
# 2) ChapterCheckpointSaver — the crash-safety core
# ═══════════════════════════════════════════════════════════════════════════


class TestChapterCheckpointSaver:
    def test_immediate_save_after_each_chapter(self, stub_env):
        """Chapters 1-2 done, crash before 3 → 1-2 fully on disk."""
        from core.pipeline.chapter_checkpoint import ChapterCheckpointSaver

        root, conv, chapters, positions = stub_env
        saver = ChapterCheckpointSaver(conv, positions, 3, build_docx=True, book_title="MyBook")
        saver.on_chapter_done(1, _mk_summary("Chapter One", 5), chapters["Chapter One"])
        saver.on_chapter_done(2, _mk_summary("Chapter Two", 4), chapters["Chapter Two"])

        assert (root / "summaries" / "001_Chapter_One_summary.md").exists()
        assert (root / "summaries" / "002_Chapter_Two_summary.md").exists()
        assert not (root / "summaries" / "003_Chapter_Three_summary.md").exists()
        # dedicated per-chapter folders with md + docx
        for pos, head in ((1, "Chapter_One"), (2, "Chapter_Two")):
            folder = root / "checkpoints" / "chapters" / f"{pos:03d}_{head}"
            assert (folder / f"{pos:03d}_{head}_summary.md").exists()
            assert (folder / f"{pos:03d}_{head}_summary.docx").exists()
        assert saver.saved_count == 2 and saver.docx_count == 2

    def test_manifest_records_chapters(self, stub_env):
        from core.pipeline.chapter_checkpoint import ChapterCheckpointSaver

        root, conv, chapters, positions = stub_env
        saver = ChapterCheckpointSaver(conv, positions, 3, build_docx=False, book_title="MyBook")
        saver.on_chapter_done(1, _mk_summary("Chapter One", 5), chapters["Chapter One"])
        saver.on_chapter_done(2, _mk_summary("Chapter Two", 4), chapters["Chapter Two"])

        manifest = json.loads((root / "checkpoints" / "chapter_manifest.json").read_text())
        assert manifest["book"] == "MyBook"
        assert manifest["version"] == "18.20.3.0"
        assert set(manifest["chapters"].keys()) == {"1", "2"}
        assert manifest["chapters"]["1"]["bullets"] == 5
        assert manifest["chapters"]["2"]["success"] is True

    def test_crash_recovery_resume_compatibility(self, stub_env):
        """THE critical property: chapter_resume parses the checkpoint files,
        so a re-run skips the completed chapters and only processes the rest."""
        from core.pipeline.chapter_checkpoint import ChapterCheckpointSaver
        from core.pipeline.chapter_resume import detect_processed_chapters

        root, conv, chapters, positions = stub_env
        saver = ChapterCheckpointSaver(conv, positions, 3, build_docx=False, book_title="MyBook")
        saver.on_chapter_done(1, _mk_summary("Chapter One", 5), chapters["Chapter One"])
        saver.on_chapter_done(2, _mk_summary("Chapter Two", 4), chapters["Chapter Two"])
        # CRASH — nothing else runs. Detect what survived:
        processed = detect_processed_chapters(root)
        assert set(processed.keys()) == {"chapter_one", "chapter_two"}
        assert processed["chapter_one"].success
        assert len(processed["chapter_one"].bullet_points) == 5

    def test_write_once_untouched_guarantee(self, stub_env):
        """A completed chapter's file is NEVER rewritten — byte + mtime identical."""
        from core.pipeline.chapter_checkpoint import ChapterCheckpointSaver

        root, conv, chapters, positions = stub_env
        saver = ChapterCheckpointSaver(conv, positions, 3, build_docx=False, book_title="MyBook")
        saver.on_chapter_done(1, _mk_summary("Chapter One", 5), chapters["Chapter One"])

        f = root / "summaries" / "001_Chapter_One_summary.md"
        before, mtime = f.read_text(), f.stat().st_mtime_ns

        # recovery run produces a DIFFERENT summary for the same chapter
        saver2 = ChapterCheckpointSaver(conv, positions, 3, build_docx=False, book_title="MyBook")
        saver2.on_chapter_done(1, _mk_summary("Chapter One", 99), chapters["Chapter One"])
        assert f.read_text() == before
        assert f.stat().st_mtime_ns == mtime
        assert saver2.skipped_existing == 1

    def test_failed_chapter_not_saved(self, stub_env):
        from core.pipeline.chapter_checkpoint import ChapterCheckpointSaver

        root, conv, chapters, positions = stub_env
        saver = ChapterCheckpointSaver(conv, positions, 3, build_docx=True, book_title="MyBook")
        failed = _mk_summary("Chapter Three", 0, success=False)
        saver.on_chapter_done(3, failed, chapters["Chapter Three"])
        assert not (root / "summaries" / "003_Chapter_Three_summary.md").exists()
        assert not (root / "checkpoints" / "chapters" / "003_Chapter_Three").exists()
        assert saver.saved_count == 0

    def test_empty_summary_not_saved(self, stub_env):
        from core.models import SummaryResult
        from core.pipeline.chapter_checkpoint import ChapterCheckpointSaver

        root, conv, chapters, positions = stub_env
        saver = ChapterCheckpointSaver(conv, positions, 3, build_docx=False, book_title="MyBook")
        empty = SummaryResult(
            chapter_heading="Chapter Three",
            executive_summary="",
            bullet_points=[],
            thorough_summary="",
            success=True,
        )
        saver.on_chapter_done(3, empty, chapters["Chapter Three"])
        assert not (root / "summaries" / "003_Chapter_Three_summary.md").exists()

    def test_best_effort_never_raises(self, stub_env):
        """Any internal failure logs a WARNING and never kills summarization."""
        from core.pipeline.chapter_checkpoint import ChapterCheckpointSaver

        root, conv, chapters, positions = stub_env

        class Boom(ChapterCheckpointSaver):
            def _update_manifest(self, *a, **k):
                raise RuntimeError("manifest boom")

        saver = Boom(conv, positions, 3, build_docx=False, book_title="MyBook")
        saver.on_chapter_done(1, _mk_summary("Chapter One", 2), chapters["Chapter One"])  # no raise
        # the md file (written before the manifest step) survived
        assert (root / "summaries" / "001_Chapter_One_summary.md").exists()

    def test_heading_position_map_overrides_index(self, stub_env):
        """Resume run: chapters 1-2 loaded, only chapter 3 processed (todo
        idx=1) → file must be numbered 003 via the heading map."""
        from core.pipeline.chapter_checkpoint import ChapterCheckpointSaver

        root, conv, chapters, positions = stub_env
        saver = ChapterCheckpointSaver(conv, positions, 3, build_docx=False, book_title="MyBook")
        saver.on_chapter_done(1, _mk_summary("Chapter Three", 2), chapters["Chapter Three"])
        assert (root / "summaries" / "003_Chapter_Three_summary.md").exists()

    def test_fallback_position_without_map(self, stub_env):
        from core.pipeline.chapter_checkpoint import ChapterCheckpointSaver

        root, conv, chapters, positions = stub_env
        saver = ChapterCheckpointSaver(conv, None, 0, build_docx=False, book_title="MyBook")
        saver.on_chapter_done(
            7, _mk_summary("Chapter Seven", 2), {"heading": "Chapter Seven", "content": "x"}
        )
        assert (root / "summaries" / "007_Chapter_Seven_summary.md").exists()

    def test_checkpoint_progress_recorded(self, stub_env):
        from core.pipeline.chapter_checkpoint import ChapterCheckpointSaver

        root, conv, chapters, positions = stub_env
        saver = ChapterCheckpointSaver(conv, positions, 3, build_docx=False, book_title="MyBook")
        saver.on_chapter_done(1, _mk_summary("Chapter One", 2), chapters["Chapter One"])
        saver.on_chapter_done(2, _mk_summary("Chapter Two", 2), chapters["Chapter Two"])
        assert conv.checkpoint.calls[-1][0] == "summarization"
        assert conv.checkpoint.calls[-1][1]["chapters_done"] == 2
        assert conv.checkpoint.calls[-1][1]["chapters_total"] == 3

    def test_per_chapter_docx_content(self, stub_env):
        from core.pipeline.chapter_checkpoint import ChapterCheckpointSaver

        root, conv, chapters, positions = stub_env
        saver = ChapterCheckpointSaver(conv, positions, 3, build_docx=True, book_title="MyBook")
        saver.on_chapter_done(1, _mk_summary("Chapter One", 3), chapters["Chapter One"])
        docx_path = (
            root / "checkpoints" / "chapters" / "001_Chapter_One" / "001_Chapter_One_summary.docx"
        )
        assert docx_path.exists()
        from docx import Document

        doc = Document(str(docx_path))
        # V0.06 (rig audit F3): extract BOTH body paragraphs AND table-cell
        # text — the per-chapter "Key Takeaways" callout is a styled TABLE
        # (python-docx `doc.paragraphs` does not include table cells), and
        # the old assertion only passed via the now-REMOVED duplicate
        # "Global Key Takeaways" page (every takeaway twice per document).
        text = "\n".join(p.text for p in doc.paragraphs)
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    text += "\n" + cell.text
        assert "Chapter One" in text
        assert "Executive Summary" in text
        assert "Key Takeaways" in text
        assert "Bullet 0 about Chapter One" in text
        # V0.06 (F3): the duplicate global-takeaways page is gone.
        assert "Global Key Takeaways" not in text
        # atomic save leaves no temp files behind
        assert not list((root / "checkpoints").rglob("*.tmp"))


# ═══════════════════════════════════════════════════════════════════════════
# 3) Checkpoint.mark_progress (phase-level state)
# ═══════════════════════════════════════════════════════════════════════════


class TestCheckpointMarkProgress:
    def _cp(self, tmp_path):
        from core.pipeline.checkpoint import Checkpoint

        cp = Checkpoint(book_name="B", input_path="x")
        cp.set_output_dir(tmp_path)
        return cp

    def test_stays_pending_and_resumable(self, tmp_path):
        cp = self._cp(tmp_path)
        cp.mark_progress("summarization", chapters_done=4, chapters_total=10)
        assert cp.phases["summarization"].status == "pending"
        assert cp.should_run("summarization") is True
        assert cp.get_resume_point() is not None

    def test_metadata_recorded_and_roundtrips(self, tmp_path):
        from core.pipeline.checkpoint import Checkpoint

        cp = self._cp(tmp_path)
        cp.mark_progress("summarization", chapters_done=4, chapters_total=10)
        loaded = Checkpoint.load(tmp_path)
        assert loaded.phases["summarization"].metadata["chapters_done"] == 4
        assert loaded.phases["summarization"].status == "pending"

    def test_done_phase_not_downgraded(self, tmp_path):
        cp = self._cp(tmp_path)
        cp.mark_done("summarization", chapters=5)
        cp.mark_progress("summarization", chapters_done=3)
        assert cp.phases["summarization"].status == "done"

    def test_unknown_phase_created(self, tmp_path):
        cp = self._cp(tmp_path)
        cp.mark_progress("brand_new_phase", x=1)
        assert cp.phases["brand_new_phase"].status == "pending"


# ═══════════════════════════════════════════════════════════════════════════
# 4) summarize_chapters chapter_done_callback
# ═══════════════════════════════════════════════════════════════════════════


def _bare_summarizer(max_workers: int = 1):
    """Summarizer without __init__ (established object.__new__ pattern)."""
    from core.summarisation.summarizer import Summarizer

    s = object.__new__(Summarizer)
    s.max_workers = max_workers
    s.log = lambda msg, level="INFO": None
    s.is_local = False
    s.local_chapter_deadline = 0
    s._chapter_deadline_start = None
    return s


class TestSummarizeChaptersCallback:
    def _patch_inner(self, s, fail_heading=None):
        chapters_seen = []

        def fake_inner(heading, content, concepts=None, callouts=None, book_context=""):
            chapters_seen.append(heading)
            from core.models import SummaryResult

            ok = heading != fail_heading
            return SummaryResult(
                chapter_heading=heading,
                executive_summary=f"Exec {heading}" if ok else "",
                thorough_summary=f"Thorough {heading}" if ok else "",
                bullet_points=[f"b1 {heading}"] if ok else [],
                success=ok,
                error_msg=None if ok else "boom",
            )

        s._summarize_inner = fake_inner
        s._distill_book_context = lambda ctx, exec, heading: ""
        return chapters_seen

    def test_sequential_callback_fires_per_chapter(self):
        s = _bare_summarizer(max_workers=1)
        self._patch_inner(s)
        chapters = [{"heading": f"Chapter {i}", "content": f"content {i}"} for i in range(1, 4)]
        calls = []
        s.summarize_chapters(
            chapters,
            chapter_done_callback=lambda idx, summ, chap: calls.append(
                (idx, summ.chapter_heading, chap["heading"])
            ),
        )
        assert calls == [
            (1, "Chapter 1", "Chapter 1"),
            (2, "Chapter 2", "Chapter 2"),
            (3, "Chapter 3", "Chapter 3"),
        ]

    def test_parallel_callback_fires_per_chapter(self):
        s = _bare_summarizer(max_workers=2)
        self._patch_inner(s)
        chapters = [{"heading": f"Chapter {i}", "content": f"content {i}"} for i in range(1, 5)]
        calls = []
        results = s.summarize_chapters(
            chapters,
            chapter_done_callback=lambda idx, summ, chap: calls.append((idx, summ.chapter_heading)),
        )
        assert len(results) == 4
        assert sorted(calls) == [
            (1, "Chapter 1"),
            (2, "Chapter 2"),
            (3, "Chapter 3"),
            (4, "Chapter 4"),
        ]

    def test_callback_exception_does_not_kill_loop(self):
        s = _bare_summarizer(max_workers=1)
        self._patch_inner(s)
        chapters = [{"heading": f"Chapter {i}", "content": f"content {i}"} for i in range(1, 4)]
        calls = []

        def bad_callback(idx, summ, chap):
            if idx == 1:
                raise RuntimeError("checkpoint subsystem down")
            calls.append(idx)

        results = s.summarize_chapters(chapters, chapter_done_callback=bad_callback)
        assert len(results) == 3  # all chapters still summarized
        assert calls == [2, 3]  # later callbacks still ran

    def test_callback_receives_failed_results_too(self):
        """The saver sees failed summaries (and skips saving them itself)."""
        s = _bare_summarizer(max_workers=1)
        self._patch_inner(s, fail_heading="Chapter 2")
        chapters = [
            {"heading": "Chapter 1", "content": "c"},
            {"heading": "Chapter 2", "content": "c"},
        ]
        seen = []
        s.summarize_chapters(
            chapters, chapter_done_callback=lambda i, r, c: seen.append((i, r.success))
        )
        assert seen == [(1, True), (2, False)]

    def test_no_callback_still_works(self):
        """Backward compat: existing callers pass no callback."""
        s = _bare_summarizer(max_workers=1)
        self._patch_inner(s)
        results = s.summarize_chapters([{"heading": "A", "content": "c"}])
        assert len(results) == 1 and results[0].success


# ═══════════════════════════════════════════════════════════════════════════
# 5) Settings + wiring
# ═══════════════════════════════════════════════════════════════════════════


class TestSettingsAndWiring:
    def test_incremental_chapter_save_default_true(self):
        from utils.settings import Settings

        assert Settings().get("summarisation.incremental_chapter_save") is True

    def test_summarizer_signature_accepts_callback(self):
        import inspect

        from core.summarisation.summarizer import Summarizer

        params = inspect.signature(Summarizer.summarize_chapters).parameters
        assert "chapter_done_callback" in params

    def test_conversion_stage_passes_callback(self):
        """The conversion stage source wires the saver into summarize_chapters."""
        src = (PROJECT_ROOT / "core" / "pipeline" / "conversion" / "summarization.py").read_text()
        assert "chapter_done_callback=" in src
        assert "ChapterCheckpointSaver" in src
        assert "incremental_chapter_save" in src
        # final writer is write-once (skip existing COMPLETE files)
        assert "already saved — leaving untouched" in src

    def test_end_of_run_writer_uses_shared_renderer(self):
        """End-of-run save and checkpoint saver produce identical bytes."""
        src = (PROJECT_ROOT / "core" / "pipeline" / "conversion" / "summarization.py").read_text()
        assert "render_summary_markdown" in src
