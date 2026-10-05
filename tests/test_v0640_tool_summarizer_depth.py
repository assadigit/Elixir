"""V0.6.4 — the user-reported Summarizer Tool backlog, defect-driven.

Six MEDIUM requirements (user report, 6 items):
1. **Multiple output formats (checkbox)**: the single Format combo → a
   Markdown / Word / PDF checkbox row; every checked format gets its own
   output file from ONE summarisation run.
2. **Custom title**: a toggle-hidden input for the document heading (e.g.
   "BMW Company History") replacing the auto ``"{stem} — {depth} summary"``
   title when set.
3. **Black output styling**: the docx + pdf outputs must be black by
   default (was: green theme for pdf, Word-blue headings for docx).
4. **Farsi text**: paste Farsi → summary in Farsi (the tool summarizes,
   never translates). Root cause: the old prompt was 100 % English with an
   empty ``{{language}}`` variable — every depth prompt now carries the
   "Write in the same language as the text." rule.
5. **Fine-tuned depth prompts**: the user-supplied Fast / Standard /
   Thorough prompt texts, adopted VERBATIM (chat-mangled line breaks
   re-flowed, wording untouched).
6. **Term change**: "Quality" → "Depth" in the tool UI.

Layers (house UI-test pattern, cf. test_v0600_tool_summarization_e2e):
1. Pure: ``core.summarisation.tool_prompts`` (verbatim prompt pins, the
   render contract, heading normalisation) + the black-theme formatters.
2. Service-level: ``_summarize_work`` with a fake ``_call_llm`` —
   multi-format writes, custom title, depth routing, Farsi round-trip.
3. Dialog-level drive: the checkbox row, the toggle-hidden title input,
   the guards, ``_on_try_another`` reset.
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest

_project_root = Path(__file__).resolve().parent.parent
if str(_project_root) not in sys.path:
    sys.path.insert(0, str(_project_root))


# ── Layer 0: availability gates ─────────────────────────────────────────────


def _qt_available() -> bool:
    try:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication  # noqa: F401

        return True
    except Exception:
        return False


def _docx_available() -> bool:
    try:
        import docx  # noqa: F401

        return True
    except Exception:
        return False


def _reportlab_available() -> bool:
    try:
        import reportlab  # noqa: F401

        return True
    except Exception:
        return False


# Content-agnostic two-paragraph input (Rule 1: no book/author/domain text).
TEXT = (
    "Bees pollinate a large share of flowering plants, transferring pollen "
    "as they forage for nectar. Their activity supports both wild ecosystems "
    "and agricultural crops.\n\n"
    "Honeybee colonies divide labor among a queen, workers, and drones. "
    "Workers perform successive roles during their lifetime, from nursing "
    "brood to foraging in the field, and communicate food locations through "
    "dance movements."
)

# A depth-style LLM response (the new prompts produce plain sections, not
# XML): a Fast-shaped answer with a bare section label + bullets.
CANNED_DEPTH = (
    "The material explains how one insect species supports plant "
    "reproduction and how its colonies organise their labor.\n\n"
    "Key Points:\n"
    "- Pollen is transferred while the insects forage for nectar.\n"
    "- Colony labor is divided among a queen, workers, and drones.\n"
    "- Workers change roles over their lifetime, ending as foragers."
)

# A Farsi (Persian) depth-style response — Farsi in must come Farsi out
# (requirement 4: the tool must round-trip the language byte-exactly).
CANNED_FARSI = (
    "این متن توضیح می‌دهد که یک گونه حشره چگونه به تولید مثل گیاهان کمک می‌کند.\n\n"
    "Key Points:\n"
    "- گرده در حین جستجوی شهد منتقل می‌شود.\n"
    "- تقسیم کار بین ملکه، کارگرها و نرها انجام می‌شود."
)

CUSTOM_TITLE = "BMW Company History"


# ── Layer 1: tool_prompts — the verbatim depth prompt bank ───────────────────


class TestToolPromptBank:
    """``core.summarisation.tool_prompts`` — the user-supplied prompts must
    ship VERBATIM (only the chat-mangled line breaks were re-flowed)."""

    def _prompts(self):
        from core.summarisation.tool_prompts import TOOL_PROMPTS

        return TOOL_PROMPTS

    def test_all_three_depths_present(self):
        prompts = self._prompts()
        assert set(prompts) == {"fast", "standard", "thorough"}

    def test_fast_prompt_verbatim_markers(self):
        p = self._prompts()["fast"]
        assert "You are a summarizer." in p
        assert "Your only job is to shorten the text you are given." in p
        assert "You never write new content." in p
        # STRICT RULES
        assert "Use ONLY information that is written in the text." in p
        assert "If something is unclear or missing in the text, leave it out. Never guess." in p
        assert "Keep names, numbers, dates, and technical terms exactly as they appear." in p
        assert "Write in the same language as the text." in p
        assert "Do not write an introduction or a closing sentence." in p
        # FORMAT
        assert "Maximum 10 bullets." in p
        assert "Each bullet must be one or two short lines." in p
        # LENGTH
        assert "Aim for about 200 words." in p
        assert "NEVER go above 300 words" in p
        assert "The summary must never be longer than the original text." in p
        assert "Skip minor examples, side stories, and small details." in p

    def test_standard_prompt_verbatim_markers(self):
        p = self._prompts()["standard"]
        assert "Your only job is to condense and organize the text you are given." in p
        assert "Do not add facts, examples, opinions, advice, conclusions, or explanations" in p
        assert (
            "Do not judge or evaluate the text. Report what it says, not what you think about it."
            in p
        )
        assert "Write in the same language as the text." in p
        # FORMAT
        assert '"Overview": 2 to 3 sentences' in p
        assert '"Key Points": a bullet list of the main ideas, in the same order they appear' in p
        assert '"Important Details": specific facts, numbers, names, or steps' in p
        assert "together must have a maximum of 20 bullets." in p
        # LENGTH
        assert "NEVER go above 500 words" in p
        assert "Remove repetition, filler, and minor examples." in p

    def test_thorough_prompt_verbatim_markers(self):
        p = self._prompts()["thorough"]
        assert (
            "Your only job is to condense, organize, and structure the text you are "
            "given so that nothing important is lost."
        ) in p
        assert "Never guess or fill gaps." in p
        assert (
            "Keep names, numbers, dates, quotes, and technical terms exactly as they appear." in p
        )
        assert (
            "If the author gives a definition, a step-by-step process, a list, or a comparison, keep its structure."
            in p
        )
        # FORMAT
        assert "Write the summary as dense bullet points." in p
        assert "There is no limit on the number of bullets." in p
        assert '"Overview": 3 to 4 sentences' in p
        assert '"Detailed Summary": follow the text\'s own structure.' in p
        assert (
            '"Key Facts and Figures": a bullet list of all important numbers, dates, names, definitions, and terms'
            in p
        )
        assert (
            '"Conclusions and Takeaways": the conclusions, recommendations, or final points that the AUTHOR states.'
            in p
        )
        assert 'write "None stated in the text." Do not create your own.' in p
        # LENGTH
        assert "NEVER be longer than 60 percent of the original text." in p
        assert "Remove only repetition and filler." in p

    def test_every_depth_carries_the_same_language_rule(self):
        """Requirement 4 (Farsi): every depth prompt must instruct the model
        to answer in the source language — the old English-only prompt bank
        (with an empty ``{{language}}``) is what translated Farsi → English."""
        for depth, p in self._prompts().items():
            assert "Write in the same language as the text." in p, depth


class TestRenderToolPrompt:
    """``render_tool_prompt`` — the dialog's prompt builder entry point."""

    def _render(self, depth, text=TEXT, model=""):
        from core.summarisation.tool_prompts import render_tool_prompt

        return render_tool_prompt(depth, text, model=model)

    def test_prompt_carries_instructions_then_text(self):
        prompt, max_tokens = self._render("fast")
        assert "You are a summarizer." in prompt
        assert TEXT[:60] in prompt
        # The instructions come BEFORE the text (a summariser contract).
        assert prompt.index("You are a summarizer.") < prompt.index(TEXT[:60])

    def test_text_is_untrusted_fenced(self):
        # V18.2.21 prompt-injection defense must survive the V0.6.4 rewrite
        # (the old path got the fence from PromptManager.render_template).
        prompt, _ = self._render("standard")
        assert "<untrusted_source>" in prompt
        assert "</untrusted_source>" in prompt

    def test_max_tokens_mapping_unchanged(self):
        assert self._render("fast")[1] == 1024
        assert self._render("standard")[1] == 2048
        assert self._render("thorough")[1] == 4096

    def test_unknown_depth_falls_back_to_standard(self):
        from core.summarisation.tool_prompts import TOOL_PROMPTS

        prompt, _ = self._render("bogus-depth")
        assert "condense and organize" in prompt  # the standard opener
        assert "condense, organize, and structure" not in prompt  # not thorough
        assert prompt.split("\n")[0] == TOOL_PROMPTS["standard"].split("\n")[0]

    def test_low_bit_quant_model_skips_the_fence(self):
        # V18.2.29: the <untrusted_source> fence tags confuse low-bit quants
        # (the render_template precedent — raw text is SAFER for them).
        prompt, _ = self._render("fast", model="Qwen3.8-27B-UD-IQ2_S-Instruct:latest")
        assert "<untrusted_source>" not in prompt
        assert TEXT[:60] in prompt

    def test_sanitizer_failure_fails_open(self, monkeypatch):
        # Defense-in-depth (the render_template precedent): the sanitizer is
        # hardening, not a gate — a crash must never block a summarisation.
        import core.llm.prompt_sanitizer as ps

        def _boom(*a, **k):
            raise RuntimeError("sanitizer exploded")

        monkeypatch.setattr(ps, "wrap_untrusted", _boom)
        prompt, _ = self._render("standard")
        assert "You are a summarizer." in prompt
        assert TEXT[:60] in prompt  # raw text substituted, run still possible

    def test_already_fenced_text_is_not_double_wrapped(self):
        pre_fenced = f"<untrusted_source>\n{TEXT}\n</untrusted_source>"
        prompt, _ = self._render("standard", text=pre_fenced)
        assert prompt.count("<untrusted_source>") == 1


class TestNormalizeDepthHeadings:
    """``normalize_depth_headings`` — the model emits bare section labels
    (``Overview:``) per the depth FORMAT contract; the writers need markdown
    headings so docx/pdf render real Heading styles and Obsidian can map
    callouts."""

    def _norm(self):
        from core.summarisation.tool_prompts import normalize_depth_headings

        return normalize_depth_headings

    def test_known_labels_become_headings(self):
        norm = self._norm()
        text = (
            "Overview:\nThe material states a claim.\n\n"
            "Key Points:\n- First point.\n\n"
            'Important Details":\n- A number appears.\n\n'
            "Detailed Summary:\nProse here.\n\n"
            "Key Facts and Figures:\n- 42 items.\n\n"
            "Conclusions and Takeaways:\n- Final point."
        )
        out = norm(text)
        assert "## Overview" in out
        assert "## Key Points" in out
        assert "## Important Details" in out
        assert "## Detailed Summary" in out
        assert "## Key Facts and Figures" in out
        assert "## Conclusions and Takeaways" in out
        # The bodies are untouched.
        assert "- First point." in out
        assert "Prose here." in out

    def test_quoted_label_forms_normalise(self):
        norm = self._norm()
        out = norm('"Overview": 2 to 3 sentences of prose.')
        assert out.startswith("## Overview")
        assert "2 to 3 sentences of prose." in out

    def test_unknown_label_is_left_alone(self):
        # A colon line that is NOT one of the prompt's section labels is
        # ordinary prose ("Note: something") — never re-styled.
        norm = self._norm()
        out = norm("Note: this is a normal sentence with a colon.")
        assert "## Note" not in out
        assert out == "Note: this is a normal sentence with a colon."

    def test_mid_sentence_labels_untouched(self):
        norm = self._norm()
        out = norm("The overview: of the chapter appears early.")
        assert "## " not in out

    def test_idempotent(self):
        norm = self._norm()
        once = norm("Overview:\nBody text.")
        assert norm(once) == once

    def test_persian_body_never_mangled(self):
        norm = self._norm()
        out = norm(CANNED_FARSI)
        assert "می‌کند" in out  # the Farsi text survived byte-exact
        assert "## Key Points" in out  # the label still normalised
        assert "- گرده در حین جستجوی شهد منتقل می‌شود." in out


# ── Layer 1: the black-theme formatters ──────────────────────────────────────


@pytest.mark.skipif(not _docx_available(), reason="python-docx unavailable")
class TestDocxBlackHeadings:
    """Requirement 3: docx output must be black by default. Pre-fix,
    ``doc.add_heading`` used the default Word template's blue theme for
    Heading 1–4 (only the level-0 title was explicitly darkened)."""

    def _doc(self, tmp_path, title="Doc title"):
        from core.outputs.service_formatters import write_docx

        out = tmp_path / "out.docx"
        text = "## Section One\n\nFirst body paragraph.\n\n### Sub Section\n\n- a bullet\n"
        assert write_docx(text, out, title=title) is True
        from docx import Document

        return Document(str(out))

    def test_every_heading_run_is_explicitly_black(self, tmp_path):
        doc = self._doc(tmp_path)
        headings = [
            p for p in doc.paragraphs if p.style.name.lower().startswith(("title", "heading"))
        ]
        assert len(headings) >= 3  # the title + H1 + H2 from the fixture
        for p in headings:
            for run in p.runs:
                color = run.font.color
                assert color is not None and color.rgb is not None, (
                    f"{p.style.name} run has no explicit color (falls back to Word blue)"
                )
                assert str(color.rgb) in ("1A1A1A", "000000"), (
                    f"{p.style.name} run color {color.rgb} is not black"
                )

    def test_heading_text_survives(self, tmp_path):
        doc = self._doc(tmp_path)
        texts = [p.text for p in doc.paragraphs]
        assert "Doc title" in texts
        assert "Section One" in texts
        assert "Sub Section" in texts


@pytest.mark.skipif(not _reportlab_available(), reason="reportlab unavailable")
class TestPdfBlackTheme:
    """Requirement 3: pdf output must be black by default (was a green
    theme: #166534/#15803d/#16a34a in BOTH the WeasyPrint CSS and the
    ReportLab styles)."""

    def _write(self, tmp_path, title="Pdf title"):
        from core.outputs.service_formatters import write_pdf

        out = tmp_path / "out.pdf"
        text = "## Section One\n\nBody paragraph.\n\n- a bullet\n"
        assert write_pdf(text, out, title=title) is True
        return out

    def test_reportlab_path_uses_black_text_colors(self, tmp_path, monkeypatch):
        # Force the ReportLab backend (simulate WeasyPrint absent — the
        # documented bare-install degradation path).
        import builtins

        real_import = builtins.__import__

        def _no_weasy(name, *a, **k):
            if "weasyprint" in name:
                raise ImportError(f"blocked for test: {name}")
            return real_import(name, *a, **k)

        monkeypatch.setattr(builtins, "__import__", _no_weasy)
        out = self._write(tmp_path)
        data = out.read_bytes()
        assert data[:5] == b"%PDF-"
        # ReportLab compresses the content stream — decompress via pypdf
        # and inspect the colour operands: near-black 0x1A/255 = .101961
        # (ReportLab omits the leading zero) for the title/H1 + the title
        # rule. The old green theme (#166534 → .086275 .396078 .203922)
        # must be gone.
        from pypdf import PdfReader

        stream = PdfReader(str(out)).pages[0].get_contents().get_data().decode("latin-1")
        assert ".101961 .101961 .101961 rg" in stream  # black title/heading text
        assert ".396078" not in stream  # the old green G operand
        assert ".203922" not in stream  # the old green B operand

    def test_source_has_no_green_theme_literals(self):
        # Belt-and-braces source pin (the v1820412 precedent — comments
        # excluded, only executable code checked): the green palette must
        # not survive anywhere in the formatter module.
        from core.outputs import service_formatters as sf

        src = Path(sf.__file__).read_text(encoding="utf-8")
        code = "\n".join(ln for ln in src.splitlines() if not ln.lstrip().startswith("#"))
        for green in ("166534", "15803d", "16a34a", "86efac", "f0fdf4"):
            assert green not in code, f"green theme literal #{green} still present"
        assert "1a1a1a" in code  # the black default is declared

    def test_pdf_write_still_succeeds_with_weasyprint(self, tmp_path):
        out = self._write(tmp_path)
        assert out.read_bytes()[:5] == b"%PDF-"

    def test_weasyprint_css_declares_black_headings(self, tmp_path):
        # Drive the WeasyPrint path (installed on this rig) and inspect the
        # module-level CSS constant it embeds.
        from core.outputs.service_formatters import _PDF_CSS

        css = _PDF_CSS
        assert "#1a1a1a" in css
        for green in ("#166534", "#15803d", "#16a34a", "#86efac", "#f0fdf4"):
            assert green not in css, f"green literal {green} in the pdf CSS"


@pytest.mark.skipif(not _docx_available(), reason="python-docx unavailable")
class TestDocxBlackRegression:
    """The title (level 0) was ALREADY black pre-fix — a characterization
    that must survive the heading fix."""

    def test_title_run_still_dark(self, tmp_path):
        from core.outputs.service_formatters import write_docx

        out = tmp_path / "t.docx"
        write_docx("Body only.", out, title="Only Title")
        from docx import Document

        doc = Document(str(out))
        title_paras = [p for p in doc.paragraphs if p.style.name.lower() == "title"]
        assert title_paras and title_paras[0].text == "Only Title"
        for run in title_paras[0].runs:
            assert str(run.font.color.rgb) == "1A1A1A"


# ── Layer 2: service-level (_summarize_work, fake LLM) ───────────────────────


# Reuse the v0600 hermetic scaffolding (one import — the house dedup).
from tests.test_v0600_tool_summarization_e2e import (  # noqa: E402
    _DIALOGS,
    _fake_llm,
    _install_settings,
    _make_dialog,
    _spin_until_done,
    _TestSettings,
)


@pytest.fixture(autouse=True)
def _v0640_join_workers_and_clear_settings():
    """The v0600 containment fixtures are module-scoped — this module gets
    its own copy: join any SimpleWorker before its dialog is collected
    (the QThread-destroyed-while-running segfault class), neutralize the
    close-guard, and clear the Settings._test_override DI hook."""
    yield
    from PySide6.QtWidgets import QApplication

    from utils.settings import Settings

    app = QApplication.instance()
    for dlg in _DIALOGS:
        worker = getattr(dlg, "_worker", None)
        if worker is not None and hasattr(worker, "wait"):
            try:
                worker.wait(15000)
            except Exception:
                pass
        try:
            dlg._worker = None
            dlg.close()
            dlg.deleteLater()
        except Exception:
            pass
    _DIALOGS.clear()
    if app is not None:
        for _ in range(5):
            app.processEvents()
    if isinstance(getattr(Settings, "_test_override", None), _TestSettings):
        Settings._test_override = None  # type: ignore[assignment]


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
class TestSummarizeWorkMultiFormat:
    """Requirement 1: every checked format gets its own output file, from a
    SINGLE summarisation run per input (one LLM call, N writes)."""

    _VALUES = {"summarisation.model": "fake-model", "summarisation.provider": "llamacpp"}

    def _work(self, tmp_path, exts, out_path=None, llm=None, calls=None, depth="standard", **kw):
        _install_settings({**self._VALUES, **(kw.pop("values", {}) or {})})
        dlg, _app = _make_dialog()
        if llm is None:
            llm = _fake_llm(canned=CANNED_DEPTH, calls=calls)
        dlg._call_llm = llm
        return dlg._summarize_work(
            [("pasted", TEXT)],
            False,
            None,
            out_path or str(tmp_path / "summary.md"),
            "fake-model",
            "__default__",
            depth,
            exts,
            "",
            False,
            0.6,
        )

    def test_md_and_docx_both_written_single_llm_call(self, tmp_path):
        calls: list = []
        out = tmp_path / "summary.md"
        result = self._work(tmp_path, exts=[".md", ".docx"], out_path=str(out), calls=calls)
        assert len(calls) == 1, "multi-format must NOT re-run the LLM per format"
        assert (tmp_path / "summary.md").exists()
        assert (tmp_path / "summary.docx").exists()
        assert all(r.startswith("✅") for r in result["results"])
        assert len(result["results"]) == 2  # one line per written file

    def test_all_three_formats_written(self, tmp_path):
        out = tmp_path / "trio.md"
        self._work(tmp_path, exts=[".md", ".docx", ".pdf"], out_path=str(out))
        assert (tmp_path / "trio.md").exists()
        assert (tmp_path / "trio.docx").exists()
        assert (tmp_path / "trio.pdf").exists()
        assert (tmp_path / "trio.pdf").read_bytes()[:5] == b"%PDF-"

    def test_single_format_keeps_exact_path(self, tmp_path):
        out = tmp_path / "exact.md"
        result = self._work(tmp_path, exts=[".md"], out_path=str(out))
        assert out.exists()
        assert result["file_arg"] == str(out)

    def test_user_typed_extension_is_only_the_base_name(self, tmp_path):
        # The checkboxes are authoritative: typing "notes.txt" with
        # .md+.docx checked writes notes.md + notes.docx (the typed suffix
        # is replaced per format — never docx bytes into a .md file).
        out = tmp_path / "notes.txt"
        self._work(tmp_path, exts=[".md", ".docx"], out_path=str(out))
        assert (tmp_path / "notes.md").exists()
        assert (tmp_path / "notes.docx").exists()
        assert not out.exists()

    def test_multi_file_mode_writes_every_format_per_file(self, tmp_path):
        _install_settings(self._VALUES)
        dlg, _app = _make_dialog()
        dlg._call_llm = _fake_llm(canned=CANNED_DEPTH)
        folder = tmp_path / "batch"
        folder.mkdir()
        result = dlg._summarize_work(
            [("fileA", TEXT), ("fileB", TEXT)],
            True,
            str(folder),
            None,
            "fake-model",
            "__default__",
            "fast",
            [".md", ".docx"],
            "",
            False,
            0.6,
        )
        assert result["total"] == 2
        for stem in ("fileA", "fileB"):
            assert (folder / f"{stem}_fast_summary.md").exists()
            assert (folder / f"{stem}_fast_summary.docx").exists()
        assert result["file_arg"] == str(folder)

    def test_exts_order_is_canonical_md_docx_pdf(self):
        # The dialog derives the ext list in a FIXED order so the output
        # naming is deterministic regardless of checkbox toggle order.

        dlg, _app = _make_dialog()
        dlg.fmt_md_chk.setChecked(True)
        dlg.fmt_pdf_chk.setChecked(True)
        dlg.fmt_docx_chk.setChecked(True)
        assert dlg._selected_exts() == [".md", ".docx", ".pdf"]
        dlg.fmt_docx_chk.setChecked(False)
        assert dlg._selected_exts() == [".md", ".pdf"]


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
class TestSummarizeWorkCustomTitle:
    """Requirement 2: the custom title replaces the auto
    ``"{stem} — {depth} summary"`` heading in every format."""

    _VALUES = {"summarisation.model": "fake-model", "summarisation.provider": "llamacpp"}

    def _work(self, tmp_path, custom_title, ext=".md", obsidian=False, depth="standard"):
        _install_settings(self._VALUES)
        dlg, _app = _make_dialog()
        dlg._call_llm = _fake_llm(canned=CANNED_DEPTH)
        out = tmp_path / f"t{ext}"
        result = dlg._summarize_work(
            [("pasted", TEXT)],
            False,
            None,
            str(out),
            "fake-model",
            "__default__",
            depth,
            [ext],
            custom_title,
            obsidian,
            0.6,
        )
        return out, result

    def test_custom_title_heads_the_markdown(self, tmp_path):
        out, _r = self._work(tmp_path, CUSTOM_TITLE)
        content = out.read_text(encoding="utf-8")
        assert content.startswith(f"# {CUSTOM_TITLE}\n")
        assert "Key Points" in content  # body intact

    def test_custom_title_heads_the_docx(self, tmp_path):
        out, _r = self._work(tmp_path, CUSTOM_TITLE, ext=".docx")
        from docx import Document

        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs]
        assert CUSTOM_TITLE in texts
        assert "pasted — standard summary" not in " ".join(texts)

    def test_custom_title_heads_the_pdf(self, tmp_path):
        out, _r = self._work(tmp_path, CUSTOM_TITLE, ext=".pdf")
        # ReportLab or WeasyPrint both render the title as a Paragraph —
        # verify the bytes exist in the extracted text.
        data = out.read_bytes()
        assert data[:5] == b"%PDF-"
        try:
            from pypdf import PdfReader

            txt = " ".join((pg.extract_text() or "") for pg in PdfReader(str(out)).pages)
            assert "BMW Company History" in txt
        except ImportError:
            pass  # pypdf absent — the %PDF magic + size checks suffice

    def test_custom_title_in_obsidian_frontmatter(self, tmp_path):
        out, _r = self._work(tmp_path, CUSTOM_TITLE, obsidian=True)
        content = out.read_text(encoding="utf-8")
        assert content.startswith("---")
        assert f'title: "{CUSTOM_TITLE}"' in content

    def test_auto_title_when_no_custom_title(self, tmp_path):
        # No custom title → the documented auto title is unchanged.
        out, _r = self._work(tmp_path, "", ext=".docx")
        from docx import Document

        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs]
        assert "pasted — standard summary" in texts

    def test_bare_md_keeps_no_title_without_custom(self, tmp_path):
        # Characterization: plain .md never carried a heading — that must
        # not change when no custom title is set.
        out, _r = self._work(tmp_path, "")
        content = out.read_text(encoding="utf-8")
        assert not content.startswith("# ")

    def test_depth_term_flows_into_auto_title(self, tmp_path):
        out, _r = self._work(tmp_path, "", ext=".docx", depth="thorough")
        from docx import Document

        doc = Document(str(out))
        texts = [p.text for p in doc.paragraphs]
        assert "pasted — thorough summary" in texts


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
class TestSummarizeWorkDepthPromptsAndLanguage:
    """Requirements 4 + 5: the depth chosen in the UI must select the
    matching fine-tuned prompt, and the model output must be written
    VERBATIM (Farsi in → Farsi out — the tool never translates)."""

    _VALUES = {"summarisation.model": "fake-model", "summarisation.provider": "llamacpp"}

    def _prompt_for(self, depth, tmp_path):
        calls: list = []
        _install_settings(self._VALUES)
        dlg, _app = _make_dialog()
        dlg._call_llm = _fake_llm(canned=CANNED_DEPTH, calls=calls)
        out = tmp_path / "o.md"
        dlg._summarize_work(
            [("pasted", TEXT)],
            False,
            None,
            str(out),
            "fake-model",
            "__default__",
            depth,
            [".md"],
            "",
            False,
            0.6,
        )
        assert len(calls) == 1
        return calls[0][1]  # the prompt

    def test_fast_depth_selects_the_fast_prompt(self, tmp_path):
        prompt = self._prompt_for("fast", tmp_path)
        assert "Your only job is to shorten the text you are given." in prompt
        assert "Maximum 10 bullets." in prompt
        assert "about 200 words" in prompt

    def test_thorough_depth_selects_the_thorough_prompt(self, tmp_path):
        prompt = self._prompt_for("thorough", tmp_path)
        assert "condense, organize, and structure" in prompt
        assert '"Conclusions and Takeaways"' in prompt

    def test_progress_line_says_depth_not_quality(self, tmp_path):
        # Requirement 6: the run banner must read "depth:", not "quality:".
        msgs: list[tuple] = []

        def _report(pct, msg):
            msgs.append((pct, str(msg)))

        _install_settings(self._VALUES)
        dlg, _app = _make_dialog()
        dlg._call_llm = _fake_llm(canned=CANNED_DEPTH)
        out = tmp_path / "banner.md"
        dlg._summarize_work(
            [("pasted", TEXT)],
            False,
            None,
            str(out),
            "fake-model",
            "__default__",
            "fast",
            [".md"],
            "",
            False,
            0.6,
            report=_report,
        )
        banner = next((m for _p, m in msgs if "Summarizing" in m), "")
        assert "depth: fast" in banner
        assert "quality:" not in banner

    def test_farsi_response_round_trips_verbatim(self, tmp_path):
        # Requirement 4 (the write side): a Farsi LLM response is written
        # byte-exact — no translation, no transliteration, no mangling.
        _install_settings(self._VALUES)
        dlg, _app = _make_dialog()
        dlg._call_llm = _fake_llm(canned=CANNED_FARSI)
        out = tmp_path / "farsi.md"
        dlg._summarize_work(
            [("pasted", TEXT)],
            False,
            None,
            str(out),
            "fake-model",
            "__default__",
            "standard",
            [".md"],
            "",
            False,
            0.6,
        )
        content = out.read_text(encoding="utf-8")
        assert "این متن توضیح می‌دهد" in content
        assert "- گرده در حین جستجوی شهد منتقل می‌شود." in content

    def test_raw_response_written_without_xml_reassembly(self, tmp_path):
        # The old path re-assembled the response into
        # "## Executive Summary / ## Key Takeaways / ## Detailed Summary"
        # from XML sections; the depth prompts' plain output must land
        # as-is (with label normalisation only).
        _install_settings(self._VALUES)
        dlg, _app = _make_dialog()
        dlg._call_llm = _fake_llm(canned=CANNED_DEPTH)
        out = tmp_path / "raw.md"
        dlg._summarize_work(
            [("pasted", TEXT)],
            False,
            None,
            str(out),
            "fake-model",
            "__default__",
            "standard",
            [".md"],
            "",
            False,
            0.6,
        )
        content = out.read_text(encoding="utf-8")
        assert "## Executive Summary" not in content
        assert "## Key Takeaways" not in content
        assert "## Key Points" in content  # the label → heading normalisation
        assert "- Pollen is transferred" in content


# ── Layer 3: dialog-level drive ──────────────────────────────────────────────


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
class TestDialogFormatCheckboxes:
    """Requirement 1 (the UI half): the Format combo is replaced by a
    three-checkbox row; the Obsidian toggle follows the .md checkbox."""

    def test_checkboxes_exist_and_default_follows_settings(self):
        _install_settings({"quick_services.defaults.summarize_format": ".md"})
        dlg, _app = _make_dialog()
        assert hasattr(dlg, "fmt_md_chk") and hasattr(dlg, "fmt_docx_chk")
        assert hasattr(dlg, "fmt_pdf_chk")
        assert dlg.fmt_md_chk.isChecked() is True
        assert dlg.fmt_docx_chk.isChecked() is False
        assert dlg.fmt_pdf_chk.isChecked() is False
        # The old single-choice combo is gone.
        assert not hasattr(dlg, "fmt_combo")

    def test_settings_default_docx_only_checks_docx(self):
        _install_settings({"quick_services.defaults.summarize_format": ".docx"})
        dlg, _app = _make_dialog()
        assert dlg.fmt_docx_chk.isChecked() is True
        assert dlg.fmt_md_chk.isChecked() is False
        assert dlg.fmt_pdf_chk.isChecked() is False

    def test_guard_when_no_format_checked(self, tmp_path, monkeypatch):
        import ui_qt.views.service_dialogs.dialogs.summarization_service_dialog as summ_mod

        recorded: list = []
        monkeypatch.setattr(summ_mod, "_show_success_modal", lambda *a, **k: recorded.append(1))
        _install_settings(
            {"summarisation.model": "fake-model", "summarisation.provider": "llamacpp"}
        )
        dlg, _app = _make_dialog()
        dlg._call_llm = _fake_llm(canned=CANNED_DEPTH)
        dlg.fmt_md_chk.setChecked(False)  # uncheck everything
        dlg.text_input.setPlainText(TEXT)
        dlg.out_edit.setText(str(tmp_path / "o.md"))
        dlg._do_summarize()
        assert dlg.status_label.text().startswith("⚠️")
        assert "format" in dlg.status_label.text().lower()
        assert dlg._worker is None  # no worker spawned
        assert not recorded  # no success modal

    def test_dialog_drive_two_formats_lands_both_files(self, tmp_path, monkeypatch):
        import ui_qt.views.service_dialogs.dialogs.summarization_service_dialog as summ_mod

        recorded: list = []
        monkeypatch.setattr(
            summ_mod,
            "_show_success_modal",
            lambda parent, title, message, file_path=None, on_try_another=None: recorded.append(
                (title, message, file_path)
            ),
        )
        _install_settings(
            {"summarisation.model": "fake-model", "summarisation.provider": "llamacpp"}
        )
        dlg, app = _make_dialog()
        dlg._call_llm = _fake_llm(canned=CANNED_DEPTH)
        dlg.text_input.setPlainText(TEXT)
        out = tmp_path / "driven.md"
        dlg.out_edit.setText(str(out))
        dlg.fmt_docx_chk.setChecked(True)
        dlg._do_summarize()
        assert _spin_until_done(app, dlg), dlg.status_label.text()
        assert dlg.status_label.text().startswith("✅")
        assert (tmp_path / "driven.md").exists()
        assert (tmp_path / "driven.docx").exists()
        assert recorded and recorded[0][0] == "Σ Summarization Complete"
        assert dlg.action_btn.text() == "Close"

    def test_obsidian_enabled_follows_md_checkbox(self):
        _install_settings({"quick_services.defaults.summarize_format": ".md"})
        dlg, _app = _make_dialog()
        assert dlg.obsidian_chk.isEnabled() is True
        dlg.fmt_md_chk.setChecked(False)
        assert dlg.obsidian_chk.isEnabled() is False
        dlg.fmt_md_chk.setChecked(True)
        assert dlg.obsidian_chk.isEnabled() is True


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
class TestDialogCustomTitleUi:
    """Requirement 2 (the UI half): a toggle reveals a hidden title input;
    the title drives the outputs; ``_on_try_another`` resets it."""

    def _driven(self, tmp_path, monkeypatch, title=None):
        import ui_qt.views.service_dialogs.dialogs.summarization_service_dialog as summ_mod

        recorded: list = []
        monkeypatch.setattr(
            summ_mod,
            "_show_success_modal",
            lambda parent, title, message, file_path=None, on_try_another=None: recorded.append(
                (title, message, file_path)
            ),
        )
        _install_settings(
            {"summarisation.model": "fake-model", "summarisation.provider": "llamacpp"}
        )
        dlg, app = _make_dialog()
        dlg._call_llm = _fake_llm(canned=CANNED_DEPTH)
        dlg.text_input.setPlainText(TEXT)
        out = tmp_path / "titled.md"
        dlg.out_edit.setText(str(out))
        if title is not None:
            dlg.custom_title_chk.setChecked(True)
            dlg.custom_title_edit.setText(title)
        return dlg, app, out, recorded

    def test_title_input_hidden_until_toggled(self):
        dlg, _app = _make_dialog()
        assert dlg.custom_title_edit.isHidden() is True
        dlg.custom_title_chk.setChecked(True)
        assert dlg.custom_title_edit.isHidden() is False
        dlg.custom_title_chk.setChecked(False)
        assert dlg.custom_title_edit.isHidden() is True

    def test_untoggling_clears_the_title(self):
        dlg, _app = _make_dialog()
        dlg.custom_title_chk.setChecked(True)
        dlg.custom_title_edit.setText("Temp Title")
        dlg.custom_title_chk.setChecked(False)
        assert dlg.custom_title_edit.text() == ""

    def test_custom_title_lands_in_output(self, tmp_path, monkeypatch):
        dlg, app, out, _rec = self._driven(tmp_path, monkeypatch, title=CUSTOM_TITLE)
        dlg._do_summarize()
        assert _spin_until_done(app, dlg), dlg.status_label.text()
        content = out.read_text(encoding="utf-8")
        assert content.startswith(f"# {CUSTOM_TITLE}\n")

    def test_try_another_resets_title_toggle(self, tmp_path, monkeypatch):
        dlg, app, _out, _rec = self._driven(tmp_path, monkeypatch, title=CUSTOM_TITLE)
        dlg._do_summarize()
        assert _spin_until_done(app, dlg), dlg.status_label.text()
        dlg._on_try_another()
        assert dlg.custom_title_chk.isChecked() is False
        assert dlg.custom_title_edit.isHidden() is True
        assert dlg.custom_title_edit.text() == ""

    def test_placeholder_guides_the_example(self):
        dlg, _app = _make_dialog()
        assert "BMW Company History" in dlg.custom_title_edit.placeholderText()


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
class TestDialogDepthTerm:
    """Requirement 6: the UI says "Depth", not "Quality"."""

    def test_options_row_label_reads_depth(self):
        from PySide6.QtWidgets import QLabel

        dlg, _app = _make_dialog()
        labels = [l.text() for l in dlg.findChildren(QLabel)]
        assert any(l.startswith("Depth:") for l in labels)
        assert not any(l.startswith("Quality:") for l in labels)

    def test_depth_combo_descriptions_match_new_prompts(self):
        dlg, _app = _make_dialog()
        items = [dlg.quality_combo.itemText(i) for i in range(dlg.quality_combo.count())]
        joined = " | ".join(items)
        assert "Default (change in Settings)" in items[0] or "__default__" in joined
        assert any("~200 words" in i for i in items)  # Fast
        assert any("500" in i for i in items)  # Standard
        assert any("60%" in i for i in items)  # Thorough
        # The old section-assembly descriptions are gone.
        assert not any("exec summary + 5 keypoints" in i for i in items)

    def test_settings_defaults_row_label_reads_depth(self):
        # The Settings → Defaults section feeds the tool's "__default__".
        src = Path("ui_qt/settings/sections/defaults_section.py").read_text(encoding="utf-8")
        assert 'f.addRow("Default depth"' in src
        assert 'f.addRow("Default quality"' not in src
