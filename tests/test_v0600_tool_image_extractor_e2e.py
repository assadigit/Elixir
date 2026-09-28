"""V0.6.0 (the dedicated-tools verification campaign) — Tool 4-a: the
Image Extractor, end-to-end.

Inventory + trace (evidence-first):
- Card: ``image_extractor`` — ``ui_qt/views/tool_dashboard.py:110`` TOOL_CARDS
  row 3 (EXTRACT).
- Routing: ``ui_qt/views/convert_view.py:215`` ``_on_tool_clicked`` →
  ``ImageExtractorDialog`` (via the ``base_dialog`` facade :216).
- Dialog: ``ui_qt/views/service_dialogs/dialogs/image_extractor_dialog.py``
  (the V18.20.4.6 O-b2 split) — ``_on_process`` :262 spawns the
  ``SimpleWorker``; ``_extract_all`` :314 dispatches by extension to the
  service; ``_on_done`` :411 renders the batch summary + modal.
- Service: ``core/extract/image_extractor.py`` — ``extract_images_from_pdf``
  :56 (embedded xrefs + vector-crop regions + smart-crop + full-page
  renders), ``extract_images_from_epub`` :729, ``extract_images_from_docx``
  :789; helpers ``_classify_image`` :447, ``_auto_trim_whitespace`` :273,
  ``_merge_rects`` :709, ``_crop_vector_regions`` :501,
  ``_save_raw_image`` :44.
- Existing coverage: source-pin only (test_v182045_p4a /
  test_v182046_p5 assert the worker wiring + module split via AST) — ZERO
  functional tests. This file is the first net.

Deps: fitz (PyMuPDF) is a pyproject ``dev`` extra ONLY (pyproject.toml:59
"pymupdf>=1.24.0", comment :55-58 — "The APP keeps fitz optional …
requirements.txt is intentionally unchanged"); ``core/system/
dependency_checker.py:152`` registers it as optional. The venv carries it,
so the real paths run here; the absent-fitz contract is the dialog's
per-file ``❌`` line (dialog :339-382 try/except). PIL + numpy are core
runtime deps (requirements.txt:46,49).

Layers (house UI-test pattern, cf. test_v0600_tool_cutter_e2e):
1. Pure characterization of the classifier / trim / rect-merge helpers
   (no Qt).
2. Service-level end-to-end on REAL fixtures (fitz-built PDFs with real
   embedded PNG streams + vector drawings, real zip EPUB/DOCX containers,
   a real python-docx document with an embedded picture).
3. Dialog-level drive: the real SimpleWorker thread + queued signals +
   ``_on_done`` on the main thread (offscreen QApplication, every modal
   monkeypatched to a recorder so ``exec()``-style calls never block).
"""

from __future__ import annotations

import os
import sys
import time
import zipfile
from pathlib import Path

import pytest

_project_root = Path(__file__).resolve().parent.parent
if str(_project_root) not in sys.path:
    sys.path.insert(0, str(_project_root))


# ── Layer 0: availability gates ─────────────────────────────────────────────


def _qt_available() -> bool:
    try:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication  # noqa: F401

        return True
    except Exception:
        return False


def _fitz_available() -> bool:
    try:
        import fitz  # noqa: F401

        return True
    except Exception:
        return False


def _pil_available() -> bool:
    try:
        from PIL import Image  # noqa: F401

        return Image  # keep the linter honest
    except Exception:
        return False


def _docx_available() -> bool:
    try:
        import docx  # noqa: F401

        return True
    except Exception:
        return False


# ── Fixture builders (real documents, synthetic content only) ──────────────


def _noise_png(w: int = 200, h: int = 150, seed: int = 0) -> bytes:
    """A real PNG of seeded RGB noise (classifies as a photo)."""
    import io

    import numpy as np
    from PIL import Image

    rng = np.random.default_rng(seed)
    arr = rng.integers(0, 256, size=(h, w, 3), dtype=np.uint8)
    buf = io.BytesIO()
    Image.fromarray(arr).save(buf, "PNG")
    return buf.getvalue()


def _build_pdf_with_image(tmp_path: Path, name: str = "with_img.pdf") -> Path:
    """A real 1-page PDF carrying one real embedded PNG stream."""
    import fitz

    p = tmp_path / name
    doc = fitz.open()
    page = doc.new_page(width=400, height=300)
    page.insert_text((36, 50), "caption text")
    page.insert_image(fitz.Rect(50, 80, 350, 280), stream=_noise_png(seed=1))
    doc.save(str(p))
    doc.close()
    return p


def _build_text_only_pdf(tmp_path: Path, n_pages: int = 3) -> Path:
    """A real multi-page PDF with text and NO images."""
    import fitz

    p = tmp_path / "text_only.pdf"
    doc = fitz.open()
    for i in range(n_pages):
        page = doc.new_page(width=400, height=300)
        page.insert_text((36, 50), f"TEXT ONLY PAGE {i + 1}")
    doc.save(str(p))
    doc.close()
    return p


def _build_pdf_image_on_page2(tmp_path: Path) -> Path:
    """3 pages; the embedded image sits ONLY on page 2 (range selection)."""
    import fitz

    p = tmp_path / "img_page2.pdf"
    doc = fitz.open()
    for i in range(3):
        page = doc.new_page(width=400, height=300)
        page.insert_text((36, 50), f"P{i + 1}")
        if i == 1:
            page.insert_image(fitz.Rect(50, 80, 350, 280), stream=_noise_png(seed=2))
    doc.save(str(p))
    doc.close()
    return p


def _build_pdf_dup_image(tmp_path: Path) -> Path:
    """The SAME PNG stream placed on pages 1 and 2 (xref dedup contract)."""
    import fitz

    p = tmp_path / "dup_img.pdf"
    doc = fitz.open()
    for _ in range(2):
        page = doc.new_page(width=400, height=300)
        page.insert_image(fitz.Rect(50, 80, 350, 280), stream=_noise_png(seed=3))
    doc.save(str(p))
    doc.close()
    return p


def _build_pdf_vector(tmp_path: Path) -> Path:
    """A real PDF page of pure VECTOR drawings (no embedded images)."""
    import fitz

    p = tmp_path / "vector.pdf"
    doc = fitz.open()
    page = doc.new_page(width=400, height=300)
    page.draw_rect(
        fitz.Rect(60, 60, 340, 240), color=(0.2, 0.4, 0.9), fill=(0.8, 0.9, 1.0), width=2
    )
    for x in range(100, 340, 60):
        page.draw_line(fitz.Point(x, 70), fitz.Point(x, 230), color=(0.1, 0.2, 0.5), width=1)
    doc.save(str(p))
    doc.close()
    return p


def _build_epub(tmp_path: Path, name: str = "book.epub") -> Path:
    """A real zip-container EPUB with one raster + one SVG image."""
    p = tmp_path / name
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr("OEBPS/img/pic.png", _noise_png(seed=2))
        zf.writestr(
            "OEBPS/img/cover.svg",
            b"<svg xmlns='http://www.w3.org/2000/svg' width='100' height='100'>"
            b"<rect width='100' height='100' fill='red'/></svg>",
        )
    return p


def _build_docx_zip_media(tmp_path: Path) -> Path:
    """A real zip-container DOCX with a raster + an EMF in word/media/."""
    p = tmp_path / "media.docx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types '
            'xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
        zf.writestr("word/media/image1.png", _noise_png(seed=3))
        zf.writestr("word/media/image2.emf", b"\x01\x00\x00\x00 placeholder-emf")
    return p


def epub_zip_bytes(tmp_path: Path, name: str, member: str) -> bytes:
    with zipfile.ZipFile(tmp_path / name) as zf:
        return zf.read(member)


def _collect_logs():
    logs: list[str] = []

    def log(msg, level="INFO"):
        logs.append(f"{level}: {msg}")

    return log, logs


# ── Layer 1: pure helper characterization (no Qt) ────────────────────────────


@pytest.mark.skipif(not _pil_available(), reason="Pillow unavailable")
class TestClassifyImageCharacterization:
    """Pin the exact current semantics of the heuristic classifier."""

    def _cls(self):
        from core.extract.image_extractor import _classify_image

        return _classify_image

    def test_noise_is_photo(self):
        import io

        from PIL import Image

        assert self._cls()(Image.open(io.BytesIO(_noise_png(seed=7)))) == "photos"

    def test_wide_few_colors_flat_is_table(self):
        from PIL import Image

        assert self._cls()(Image.new("RGB", (300, 100), (200, 30, 30))) == "tables"

    def test_wide_few_colors_dense_edges_is_chart(self):
        from PIL import Image, ImageDraw

        img = Image.new("RGB", (300, 100), (240, 240, 240))
        d = ImageDraw.Draw(img)
        for x in range(0, 300, 4):
            d.line([(x, 0), (x, 100)], fill=(20, 20, 20))
        for y in range(0, 100, 4):
            d.line([(0, y), (300, y)], fill=(20, 20, 20))
        assert self._cls()(img) == "charts"

    def test_line_art_square_is_diagram(self):
        from PIL import Image, ImageDraw

        img = Image.new("RGB", (150, 150), (255, 255, 255))
        d = ImageDraw.Draw(img)
        for x in range(0, 150, 3):
            d.line([(x, 0), (x, 150)], fill=(0, 0, 0))
        assert self._cls()(img) == "diagrams"

    def test_portrait_noise_is_photo(self):
        import io

        from PIL import Image

        assert self._cls()(Image.open(io.BytesIO(_noise_png(w=100, h=200, seed=9)))) == "photos"

    def test_smooth_gradient_falls_back_to_figure(self):
        from PIL import Image

        img = Image.new("RGB", (120, 120))
        for y in range(120):
            v = int(y * 255 / 119)
            for x in range(120):
                img.putpixel((x, y), (v, v, v))
        assert self._cls()(img) == "figures"

    def test_solid_square_falls_back_to_figure(self):
        from PIL import Image

        assert self._cls()(Image.new("RGB", (120, 120), (10, 200, 10))) == "figures"


@pytest.mark.skipif(not _pil_available(), reason="Pillow unavailable")
class TestAutoTrimWhitespaceCharacterization:
    """Pin the V11.76.0 whitespace trimmer."""

    def _trim(self):
        from core.extract.image_extractor import _auto_trim_whitespace

        return _auto_trim_whitespace

    def test_white_border_is_trimmed(self):
        from PIL import Image, ImageDraw

        img = Image.new("RGB", (100, 80), (255, 255, 255))
        ImageDraw.Draw(img).rectangle([30, 20, 60, 50], fill=(255, 0, 0))
        assert self._trim()(img).size == (35, 35)  # content 31x31 + 2px pad

    def test_all_white_returned_as_is(self):
        from PIL import Image

        img = Image.new("RGB", (60, 60), (255, 255, 255))
        assert self._trim()(img).size == (60, 60)

    def test_full_bleed_not_trimmed(self):
        from PIL import Image

        img = Image.new("RGB", (60, 60), (0, 128, 255))
        assert self._trim()(img).size == (60, 60)

    def test_near_white_below_threshold_counts_as_white(self):
        from PIL import Image, ImageDraw

        img = Image.new("RGB", (60, 60), (255, 255, 255))
        ImageDraw.Draw(img).rectangle([10, 10, 20, 20], fill=(250, 250, 250))
        assert self._trim()(img).size == (60, 60)  # diff 5 <= threshold 10

    def test_pad_is_clamped_at_image_edge(self):
        from PIL import Image, ImageDraw

        img = Image.new("RGB", (60, 60), (255, 255, 255))
        ImageDraw.Draw(img).rectangle([0, 0, 20, 20], fill=(0, 0, 0))
        assert self._trim()(img).size == (23, 23)  # 21x21 content + 2px, clamped

    def test_rgba_input_is_converted(self):
        from PIL import Image

        img = Image.new("RGBA", (40, 40), (255, 0, 0, 255))
        out = self._trim()(img)
        assert out.mode == "RGB"
        assert out.size == (40, 40)


@pytest.mark.skipif(not _fitz_available(), reason="pymupdf unavailable")
class TestMergeRectsCharacterization:
    """Pin the drawing-rect clustering used by the vector crop."""

    def _merge(self):
        from core.extract.image_extractor import _merge_rects

        return _merge_rects

    def test_empty(self):
        assert self._merge()([]) == []

    def test_diagonally_adjacent_rects_merge(self):
        import fitz

        merged = self._merge()(
            [fitz.Rect(0, 0, 100, 100), fitz.Rect(115, 105, 200, 200)], threshold=30
        )
        assert [tuple(r) for r in merged] == [(0.0, 0.0, 200.0, 200.0)]

    def test_vertically_aligned_rects_merge(self):
        import fitz

        merged = self._merge()(
            [fitz.Rect(0, 0, 100, 100), fitz.Rect(5, 115, 50, 160)], threshold=30
        )
        assert len(merged) == 1
        assert tuple(merged[0]) == (0.0, 0.0, 100.0, 160.0)

    def test_far_apart_rects_stay_separate(self):
        import fitz

        merged = self._merge()(
            [fitz.Rect(0, 0, 100, 100), fitz.Rect(500, 500, 600, 600)], threshold=30
        )
        assert len(merged) == 2


# ── Layer 2: service-level end-to-end on real fixtures ──────────────────────


@pytest.mark.skipif(not _fitz_available(), reason="pymupdf unavailable")
class TestPdfExtractionEndToEnd:
    """The real fitz extraction pipeline on real PDFs."""

    def _extract(self, pdf: Path, out: Path, **kwargs):
        from core.extract.image_extractor import extract_images_from_pdf

        log, logs = _collect_logs()
        progress: list[tuple[int, str]] = []

        counts = extract_images_from_pdf(
            str(pdf),
            str(out),
            log_callback=log,
            progress_callback=lambda pct, msg: progress.append((pct, msg)),
            **kwargs,
        )
        return counts, logs, progress

    def test_embedded_image_extracted_with_full_page(self, tmp_path):
        pdf = _build_pdf_with_image(tmp_path)
        out = tmp_path / "out"
        counts, _, _ = self._extract(pdf, out)
        assert counts["photos"] == 1
        assert counts["full_pages"] == 1
        assert counts["total"] == 2

        from PIL import Image

        photo = out / "photos" / "page_0001_img_001.jpg"
        assert photo.exists()
        assert Image.open(photo).size == (200, 150)  # intrinsic pixels preserved
        full = out / "full_pages" / "page_0001.jpg"
        assert full.exists()
        assert Image.open(full).size[0] > 800  # 400pt @ 150 dpi

    def test_text_only_pdf_is_graceful_skipped(self, tmp_path):
        pdf = _build_text_only_pdf(tmp_path)
        out = tmp_path / "out"
        counts, logs, _ = self._extract(pdf, out)
        assert counts["total"] == 0
        assert all(counts[c] == 0 for c in ("photos", "charts", "tables", "diagrams", "figures"))
        assert "error" not in counts
        # The category folders are created up front…
        for cat in ("photos", "charts", "tables", "diagrams", "figures", "full_pages"):
            assert (out / cat).is_dir()
        # …but nothing is written (no phantom full-page renders of text pages).
        assert not any(out.rglob("*.jpg"))
        assert any("text-only pages skipped" in m for m in logs)

    def test_page_range_selects_pages(self, tmp_path):
        pdf = _build_pdf_image_on_page2(tmp_path)
        counts, _, _ = self._extract(pdf, tmp_path / "p2", page_range=(2, 2))
        assert counts["photos"] == 1
        counts, _, _ = self._extract(pdf, tmp_path / "p1", page_range=(1, 1))
        assert counts["total"] == 0

    def test_page_range_beyond_document_is_graceful(self, tmp_path):
        pdf = _build_pdf_image_on_page2(tmp_path)
        counts, _, _ = self._extract(pdf, tmp_path / "p99", page_range=(10, 20))
        assert counts["total"] == 0
        assert "error" not in counts

    def test_same_image_on_two_pages_extracted_once(self, tmp_path):
        pdf = _build_pdf_dup_image(tmp_path)
        out = tmp_path / "out"
        counts, _, _ = self._extract(pdf, out)
        # V11.74.0 global-xref dedup: one photo, filed under its FIRST page.
        assert counts["photos"] == 1
        assert list((out / "photos").glob("page_0001_*.jpg"))

    def test_tiny_embedded_image_skipped(self, tmp_path):
        import fitz

        p = tmp_path / "tiny.pdf"
        doc = fitz.open()
        page = doc.new_page(width=400, height=300)
        page.insert_image(fitz.Rect(100, 100, 130, 130), stream=_noise_png(w=30, h=30, seed=4))
        doc.save(str(p))
        doc.close()
        counts, _, _ = self._extract(p, tmp_path / "out")
        assert counts["total"] == 0

    def test_full_pages_disabled_writes_no_renders(self, tmp_path):
        pdf = _build_pdf_with_image(tmp_path)
        out = tmp_path / "out"
        counts, _, _ = self._extract(pdf, out, extract_full_pages=False)
        assert counts["photos"] == 1
        assert counts["full_pages"] == 0
        assert not (out / "full_pages").exists()
        assert list((out / "photos").glob("*.jpg"))

    def test_unopenable_pdf_marks_error_not_zero_success(self, tmp_path):
        """V0.6.0 (D1): an unopenable file must carry counts["error"] — the
        dialog used to render it as "✅ 0 images"."""
        pdf = tmp_path / "corrupt.pdf"
        pdf.write_bytes(b"this is not a pdf at all")
        out = tmp_path / "out"
        counts, logs, _ = self._extract(pdf, out)
        assert counts["total"] == 0
        assert counts.get("error", "").startswith("Cannot open PDF")
        assert any("Cannot open PDF" in m and "ERROR" in m for m in logs)

    def test_vector_drawing_page_crops_region_and_renders(self, tmp_path):
        pdf = _build_pdf_vector(tmp_path)
        out = tmp_path / "out"
        counts, _, _ = self._extract(pdf, out)
        # One vector-crop region + one full-page render.
        assert counts["total"] == 2
        assert counts["full_pages"] == 1
        assert sum(counts[c] for c in ("photos", "charts", "tables", "diagrams", "figures")) == 1
        vec_files = [f for f in out.rglob("page_0001_vec_*.jpg")]
        assert len(vec_files) == 1

    def test_callbacks_drive_progress_to_100(self, tmp_path):
        pdf = _build_pdf_with_image(tmp_path)
        _, logs, progress = self._extract(pdf, tmp_path / "out")
        assert progress[-1][0] == 100
        assert any("Extraction complete" in m for m in logs)


@pytest.mark.skipif(not _fitz_available(), reason="pymupdf unavailable")
class TestVectorCropTrimContract:
    """V0.6.0 (D3): the vector-crop path must SAVE the auto-trimmed image.

    Evidence: ``_crop_vector_regions`` computed the V11.76.0 trimmed image
    but wrote the raw untrimmed pixmap bytes — the trim was silently
    discarded (only the classification saw it).
    """

    def test_saved_region_is_trimmed(self, tmp_path):
        import io

        import fitz
        from PIL import Image

        from core.extract.image_extractor import (
            _auto_trim_whitespace,
            _crop_vector_regions,
        )

        p = tmp_path / "sparse.pdf"
        doc = fitz.open()
        page = doc.new_page(width=400, height=300)
        # Real ink: ONE filled square with wide white margins around it.
        page.draw_rect(fitz.Rect(60, 60, 160, 160), fill=(0.2, 0.6, 0.9), color=(0.1, 0.3, 0.5))
        doc.save(str(p))
        doc.close()

        doc = fitz.open(str(p))
        page = doc[0]
        out = tmp_path / "out"
        for cat in ("photos", "charts", "tables", "diagrams", "figures"):
            (out / cat).mkdir(parents=True)
        log, _ = _collect_logs()
        dpi = 150

        drawings = [{"rect": fitz.Rect(60, 60, 160, 160)}]
        counts = _crop_vector_regions(page, 1, out, dpi, drawings, log)
        doc.close()

        assert sum(counts.values()) == 1
        saved = next(out.rglob("page_0001_vec_*.jpg"))
        saved_img = Image.open(saved)

        # Independent expectation: render the same padded clip, trim it with
        # the module's own trimmer — the SAVED file must equal that result.
        zoom = dpi / 72.0
        doc = fitz.open(str(p))
        pix = doc[0].get_pixmap(matrix=fitz.Matrix(zoom, zoom), clip=fitz.Rect(55, 55, 165, 165))
        untrimmed = Image.open(io.BytesIO(pix.tobytes("jpeg")))
        doc.close()
        trimmed = _auto_trim_whitespace(untrimmed)
        assert trimmed.size[0] < untrimmed.size[0]  # the trim really fires here
        assert saved_img.size == trimmed.size  # saved == trimmed (was: untrimmed)


@pytest.mark.skipif(not _pil_available(), reason="Pillow unavailable")
class TestEpubExtractionEndToEnd:
    """The real zipfile-based EPUB extraction."""

    def _extract(self, epub: Path, out: Path):
        from core.extract.image_extractor import extract_images_from_epub

        log, logs = _collect_logs()
        counts = extract_images_from_epub(str(epub), str(out), log_callback=log)
        return counts, logs

    def test_raster_and_svg_both_extracted(self, tmp_path):
        """V0.6.0 (D4): the EPUB filter SELECTS .svg files, but PIL cannot
        decode them — every SVG used to be dropped with a warning. They are
        now copied through verbatim into the fallback category."""
        epub = _build_epub(tmp_path)
        out = tmp_path / "out"
        counts, logs = self._extract(epub, out)
        assert counts["photos"] == 1
        assert counts["figures"] == 1
        assert counts["total"] == 2
        svg = out / "figures" / "epub_0002.svg"
        assert svg.exists()
        assert svg.read_bytes() == epub_zip_bytes(tmp_path, "book.epub", "OEBPS/img/cover.svg")
        assert not any("cover.svg" in m and "WARNING" in m for m in logs)
        assert (out / "photos" / "epub_0001.jpg").exists()

    def test_corrupt_member_is_skipped_with_warning(self, tmp_path):
        p = tmp_path / "mixed.epub"
        with zipfile.ZipFile(p, "w") as zf:
            zf.writestr("OEBPS/good.png", _noise_png(seed=5))
            zf.writestr("OEBPS/bad.png", b"\x89PNG\r\n\x1a\n not really a png")
        counts, logs = self._extract(p, tmp_path / "out")
        assert counts["total"] == 1
        assert any("bad.png" in m and "WARNING" in m for m in logs)

    def test_not_a_zip_marks_error(self, tmp_path):
        """V0.6.0 (D1): an unopenable container carries counts["error"]."""
        p = tmp_path / "bad.epub"
        p.write_bytes(b"not a zip")
        counts, logs = self._extract(p, tmp_path / "out")
        assert counts["total"] == 0
        assert counts.get("error", "").startswith("Cannot open EPUB")
        assert any("ERROR" in m for m in logs)


@pytest.mark.skipif(not _pil_available(), reason="Pillow unavailable")
class TestDocxExtractionEndToEnd:
    """The real zipfile-based DOCX media extraction."""

    def _extract(self, docx_path: Path, out: Path):
        from core.extract.image_extractor import extract_images_from_docx

        log, logs = _collect_logs()
        counts = extract_images_from_docx(str(docx_path), str(out), log_callback=log)
        return counts, logs

    @pytest.mark.skipif(not _docx_available(), reason="python-docx unavailable")
    def test_real_docx_with_picture(self, tmp_path):
        import io

        import docx
        from PIL import Image

        p = tmp_path / "real.docx"
        d = docx.Document()
        d.add_paragraph("PIC")
        buf = io.BytesIO()
        Image.open(io.BytesIO(_noise_png(seed=3))).save(buf, "PNG")
        buf.seek(0)
        d.add_picture(buf)
        d.save(str(p))

        counts, logs = self._extract(p, tmp_path / "out")
        assert counts["photos"] == 1
        assert counts["total"] == 1
        out_file = tmp_path / "out" / "photos" / "docx_0001.jpg"
        assert out_file.exists()
        assert Image.open(out_file).size == (200, 150)
        assert any("Found 1 images" in m for m in logs)

    def test_emf_media_copied_through(self, tmp_path):
        """V0.6.0 (D4): word/media can carry EMF/WMF/SVG — PIL cannot decode
        them; they are now copied verbatim instead of dropped."""
        p = _build_docx_zip_media(tmp_path)
        out = tmp_path / "out"
        counts, logs = self._extract(p, out)
        assert counts["photos"] == 1
        assert counts["figures"] == 1
        assert counts["total"] == 2
        emf = out / "figures" / "docx_0002.emf"
        assert emf.exists()
        assert emf.read_bytes() == b"\x01\x00\x00\x00 placeholder-emf"
        assert not any("image2.emf" in m and "WARNING" in m for m in logs)

    def test_not_a_zip_marks_error(self, tmp_path):
        p = tmp_path / "bad.docx"
        p.write_bytes(b"not a zip")
        counts, _ = self._extract(p, tmp_path / "out")
        assert counts.get("error", "").startswith("Cannot open DOCX")


# ── Layer 3: dialog-level drive (the real worker thread + queued signals) ──


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
@pytest.mark.skipif(not _fitz_available(), reason="pymupdf unavailable (fixture builder)")
class TestImageExtractorDialogEndToEndDrive:
    """Drive ``_on_process`` exactly as the UI does: the SimpleWorker
    thread, the queued signals, ``_on_done`` on the main thread — every
    modal monkeypatched to a recorder (``exec()``-style modals must never
    block a headless test)."""

    def _make_dialog(self, monkeypatch, tmp_path):
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])
        import ui_qt.views.service_dialogs.dialogs.image_extractor_dialog as ied_mod

        recorded: list[tuple[str, str, str]] = []

        class _RecorderBox:
            @staticmethod
            def information(parent, title, text):
                recorded.append(("information", title, text))

            @staticmethod
            def warning(parent, title, text):
                recorded.append(("warning", title, text))

            @staticmethod
            def critical(parent, title, text):
                recorded.append(("critical", title, text))

        monkeypatch.setattr(ied_mod, "QMessageBox", _RecorderBox)
        dlg = ied_mod.ImageExtractorDialog(_StubController())
        return dlg, app, recorded

    def _spin_until_done(self, app, dlg, timeout_ms: int = 15000):
        waited = 0
        while waited < timeout_ms:
            app.processEvents()
            if dlg.status_label.text().startswith(("✅", "❌", "⚠️")):
                break
            time.sleep(0.02)
            waited += 20

    def test_full_flow_lands_success_status_and_modal(self, tmp_path, monkeypatch):
        dlg, app, recorded = self._make_dialog(monkeypatch, tmp_path)
        src = _build_pdf_with_image(tmp_path)
        out = tmp_path / "driven_out"
        dlg._files = [str(src)]
        dlg.out_edit.setText(str(out))

        dlg._on_process()
        self._spin_until_done(app, dlg)

        status = dlg.status_label.text()
        assert status.startswith("✅"), status
        assert "with_img: 2 images" in status
        assert "photos: 1" in status
        # The real file landed on disk through the worker thread.
        assert (out / "with_img" / "photos" / "page_0001_img_001.jpg").exists()
        # The modal fired (recorded, never exec'd) with the real output path.
        assert recorded and recorded[0][0] == "information"
        assert recorded[0][1] == "✅ Extraction Complete"
        assert str(out) in recorded[0][2]
        # The button is restored for another run.
        assert dlg.process_btn.isEnabled()
        assert dlg.process_btn.text() == "🖼️ Extract Images"
        assert not dlg.progress_bar.isVisible()

    def test_unopenable_file_lands_error_not_zero_success(self, tmp_path, monkeypatch):
        """V0.6.0 (D1): a corrupt PDF used to end as "✅ Done! … 0 images" —
        now it lands as a per-file ❌ line, an ❌ status and a critical
        modal."""
        dlg, app, recorded = self._make_dialog(monkeypatch, tmp_path)
        bad = tmp_path / "corrupt.pdf"
        bad.write_bytes(b"not a pdf")
        out = tmp_path / "out"
        dlg._files = [str(bad)]
        dlg.out_edit.setText(str(out))

        dlg._on_process()
        self._spin_until_done(app, dlg)

        status = dlg.status_label.text()
        assert status.startswith("❌"), status
        assert "corrupt: Cannot open PDF" in status
        assert recorded and recorded[0][0] == "critical"
        assert not any(r[0] == "information" for r in recorded)
        assert dlg.process_btn.isEnabled()

    def test_mixed_batch_lands_warning_status(self, tmp_path, monkeypatch):
        dlg, app, recorded = self._make_dialog(monkeypatch, tmp_path)
        good = _build_pdf_with_image(tmp_path)
        bad = tmp_path / "corrupt.pdf"
        bad.write_bytes(b"not a pdf")
        out = tmp_path / "out"
        dlg._files = [str(good), str(bad)]
        dlg.out_edit.setText(str(out))

        dlg._on_process()
        self._spin_until_done(app, dlg)

        status = dlg.status_label.text()
        assert status.startswith("⚠️"), status
        assert "with_img: 2 images" in status
        assert "corrupt: Cannot open PDF" in status
        # Partial success still shows the informative modal.
        assert any(r[0] == "information" for r in recorded)
        assert (out / "with_img" / "photos" / "page_0001_img_001.jpg").exists()

    def test_unsupported_extension_reports_error_line(self, tmp_path, monkeypatch):
        dlg, app, recorded = self._make_dialog(monkeypatch, tmp_path)
        txt = tmp_path / "notes.txt"
        txt.write_text("plain text")
        out = tmp_path / "out"
        dlg._files = [str(txt)]
        dlg.out_edit.setText(str(out))

        dlg._on_process()
        self._spin_until_done(app, dlg)

        status = dlg.status_label.text()
        assert status.startswith("❌"), status
        assert "unsupported format (.txt)" in status
        assert recorded and recorded[0][0] == "critical"
        assert dlg.process_btn.isEnabled()

    def test_worker_failure_lands_error_status_and_modal(self, tmp_path, monkeypatch):
        """The ``failed`` signal path: an exception inside the worker body
        must re-enable the button and surface the message."""
        dlg, app, recorded = self._make_dialog(monkeypatch, tmp_path)
        out = tmp_path / "out"
        dlg._files = [str(_build_pdf_with_image(tmp_path))]
        dlg.out_edit.setText(str(out))

        def _boom(_options):
            raise RuntimeError("boom in worker")

        monkeypatch.setattr(dlg, "_extract_all", _boom)
        dlg._on_process()
        self._spin_until_done(app, dlg)

        status = dlg.status_label.text()
        assert status.startswith("❌"), status
        assert "boom in worker" in status
        assert recorded and recorded[-1][0] == "critical"
        assert dlg.process_btn.isEnabled()
        assert dlg.process_btn.text() == "🖼️ Extract Images"

    def test_missing_inputs_guard_no_worker_spawned(self, tmp_path, monkeypatch):
        dlg, _, recorded = self._make_dialog(monkeypatch, tmp_path)
        dlg._on_process()  # no files
        assert recorded and recorded[0][0] == "warning"
        assert recorded[0][1] == "No files"
        assert dlg._worker is None

        dlg._files = [str(_build_pdf_with_image(tmp_path))]
        dlg._on_process()  # files but no output folder
        assert recorded[-1][0] == "warning"
        assert recorded[-1][1] == "No output"
        assert dlg._worker is None

    def test_inverted_page_range_guards_no_worker_spawned(self, tmp_path, monkeypatch):
        """V0.6.0 (D2): From=5, To=2 used to run the worker and report
        "✅ 0 images" (the service processes an EMPTY range) — now the
        dialog refuses the impossible range up front."""
        dlg, _, recorded = self._make_dialog(monkeypatch, tmp_path)
        dlg._files = [str(_build_pdf_with_image(tmp_path))]
        out = tmp_path / "out"
        dlg.out_edit.setText(str(out))
        dlg.cb_all_pages.setChecked(False)
        dlg.spin_from.setValue(5)
        dlg.spin_to.setValue(2)

        dlg._on_process()

        assert dlg._worker is None
        assert recorded and recorded[-1][0] == "warning"
        assert recorded[-1][1] == "Invalid page range"
        assert not out.exists()  # nothing was written
        # A valid range still runs.
        dlg.spin_to.setValue(6)
        dlg._on_process()
        assert dlg._worker is not None


# ── Stub controller (the dialog only stores it; get/update_setting kept ─────
#    for parity with the reference pattern's controller contract) ───────────


class _StubController:
    def __init__(self) -> None:
        self._settings: dict[str, object] = {}

    def get_setting(self, key: str, default=None):
        return self._settings.get(key, default)

    def update_setting(self, key: str, value) -> None:
        self._settings[key] = value
