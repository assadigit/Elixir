"""V0.6.0 (the dedicated-tools verification campaign) — Tool 4-b: the
Document → Markdown converter (``pdf_to_markdown``), end-to-end.

Inventory + trace (evidence-first):
- Card: ``pdf_to_markdown`` — ``ui_qt/views/tool_dashboard.py:59`` TOOL_CARDS
  (row 1 PROCESS; the list starts at :41).
- Routing: ``ui_qt/views/convert_view.py:152`` ``_on_tool_clicked`` →
  ``QuickConvertDialog`` (via the ``base_dialog`` facade re-export,
  ``ui_qt/views/service_dialogs/base_dialog.py:37``).
- Dialog: ``ui_qt/views/service_dialogs/dialogs/quick_convert_dialog.py``
  (the V18.20.4.6 O-b2 split) — engine combo + anydoc-absent fallback
  :164-175, ``_validate_engine_format`` :316, ``_do_convert`` :378,
  ``_convert_work`` (worker-thread callable, W-b5) :436,
  ``_on_convert_finished`` :527, ``_extract_with_markitdown`` :554.
- Service: ``core/extract/anydoc_extractor.py`` — ``ANYDOC_EXTENSIONS`` :43,
  ``MARKITDOWN_EXTENSIONS`` :72, ``is_anydoc_available`` :100,
  ``_repair_pdf_with_pikepdf`` :147, ``_is_malformed_pdf_error`` :176,
  ``_count_page_markers`` :208, ``_coerce_page_count`` :213,
  ``_try_anydoc_extract`` :224, ``extract_with_anydoc`` :299.  SWOT W8:
  ``core/extract`` 0% coverage.
- THE key contract (``requirements.txt:15-18``): ``anydoc`` was REMOVED from
  the requirements (never published on PyPI), so the shipped state MUST
  auto-fall back to the markitdown engine — verified live here (anydoc is
  not installed in this venv).

Existing coverage: source-pin only (test_v182045_p4a_service_dialog_threads.py /
test_v182046_p5_dialog_split.py assert the worker wiring + result keys via
AST) — ZERO functional tests. This file is the first net.

Layers (house UI-test pattern, cf. test_v0600_tool_cutter_e2e):
1. Pure characterization of the engine-selection / format-dispatch helpers
   (service module is Qt-free; the anydoc API dispatch is pinned with fake
   ``anydoc`` modules injected into ``sys.modules``).
2. Service-level end-to-end: real fitz-built PDF / python-docx-built DOCX /
   ebooklib-built EPUB / openpyxl XLSX / python-pptx PPTX fixtures through
   the worker-thread callable ``_convert_work`` (widget-free by design).
3. Dialog-level drive: the real SimpleWorker thread + queued signals +
   ``_on_convert_finished`` on the main thread (offscreen QApplication, the
   modal monkeypatched to a recorder so ``msg.exec()`` — shared.py:111 —
   cannot block a test).
"""

from __future__ import annotations

import importlib
import os
import sys
import types
from pathlib import Path

import pytest

_project_root = Path(__file__).resolve().parent.parent
if str(_project_root) not in sys.path:
    sys.path.insert(0, str(_project_root))


# ── Layer 0: availability gates ─────────────────────────────────────────────


def _qt_available() -> bool:
    try:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication  # noqa: F401

        return True
    except Exception:
        return False


def _fitz_available() -> bool:
    try:
        import fitz  # noqa: F401

        return True
    except Exception:
        return False


def _docx_available() -> bool:
    try:
        import docx  # noqa: F401

        return True
    except Exception:
        return False


def _epub_available() -> bool:
    try:
        from ebooklib import epub  # noqa: F401

        return True
    except Exception:
        return False


def _xlsx_available() -> bool:
    try:
        import openpyxl  # noqa: F401

        return True
    except Exception:
        return False


def _pptx_available() -> bool:
    try:
        from pptx import Presentation  # noqa: F401

        return True
    except Exception:
        return False


def _anydoc_installed() -> bool:
    from core.extract.anydoc_extractor import is_anydoc_available

    return is_anydoc_available()


def _ensure_real_markitdown() -> bool:
    """Repair a stubbed ``sys.modules["markitdown"]`` (test-infra defect).

    33 test files carry a copy-pasted ``if "markitdown" not in
    sys.modules`` stub block (the exact shape the conftest fixed for
    tiktoken — see tests/conftest.py:19-45). Whichever file is collected
    first installs a ``MarkItDown = lambda *a, **k: None`` stub that stays
    for the whole session and breaks every REAL conversion in a full-suite
    run (alphabetical collection order puts this file AFTER the stub
    installers). Restore the real package so the markitdown-engine tests
    exercise the real code paths; every later guarded stub then no-ops.
    """
    mod = sys.modules.get("markitdown")
    if mod is not None and getattr(mod, "__file__", None):
        return True  # already the real package
    try:
        if mod is not None:
            del sys.modules["markitdown"]
        importlib.import_module("markitdown")
        return True
    except Exception:
        return False


_MARKITDOWN_REAL = _ensure_real_markitdown()


# ── Layer 1A: pure service helpers (no Qt, no anydoc needed) ────────────────


class TestAnydocExtractorPureHelpers:
    """Pin the exact semantics of the Qt-free helpers in
    ``core/extract/anydoc_extractor.py`` (the module had 0% coverage)."""

    def test_count_page_markers(self):
        from core.extract.anydoc_extractor import _count_page_markers

        assert _count_page_markers("<!-- page 1 -->\ntext\n<!-- page 2 -->") == 2
        assert _count_page_markers("<!-- page  3 -->") == 1  # inner spaces
        assert _count_page_markers("") == 0

    def test_count_page_markers_ignores_non_numeric(self):
        from core.extract.anydoc_extractor import _count_page_markers

        assert _count_page_markers("<!-- page x -->\npage1\npage 2x") == 0

    def test_coerce_page_count_valid_shapes(self):
        from core.extract.anydoc_extractor import _coerce_page_count

        assert _coerce_page_count(7, "") == 7
        assert _coerce_page_count("12", "") == 12

    def test_coerce_page_count_invalid_falls_back_to_markers(self):
        from core.extract.anydoc_extractor import _coerce_page_count

        text = "a\n<!-- page 1 -->\nb\n<!-- page 2 -->"
        for bad in (0, -3, None, "junk", 2.5):
            assert _coerce_page_count(bad, text) == 2

    def test_is_malformed_pdf_error_matches_classic_shapes(self):
        from core.extract.anydoc_extractor import _is_malformed_pdf_error

        class MalformedError(Exception):
            pass

        class PDFSyntaxError(Exception):
            pass

        assert _is_malformed_pdf_error(MalformedError("bad object")) is True
        assert _is_malformed_pdf_error(PDFSyntaxError("expected xref")) is True
        assert _is_malformed_pdf_error(Exception("EOF marker not found")) is True
        assert _is_malformed_pdf_error(Exception("invalid pdf header")) is True

    def test_d1_is_malformed_pdf_error_matches_truncated_stream_shape(self):
        """V0.6.0 (D1): pypdf>=4 raises ``PdfStreamError: Stream has ended
        unexpectedly`` for truncated PDFs — the most common malformed-PDF
        shape. The heuristic previously missed it, so the pikepdf repair
        path never fired for it."""
        from pypdf.errors import PdfStreamError

        from core.extract.anydoc_extractor import _is_malformed_pdf_error

        assert _is_malformed_pdf_error(PdfStreamError("Stream has ended unexpectedly")) is True

    def test_is_malformed_pdf_error_ignores_genuine_failures(self):
        """Encrypted / permission errors must NOT trigger a repair attempt
        (repair cannot rescue them — they'd just burn a retry)."""
        from pypdf.errors import FileNotDecryptedError

        from core.extract.anydoc_extractor import _is_malformed_pdf_error

        assert (
            _is_malformed_pdf_error(FileNotDecryptedError("File has not been decrypted")) is False
        )
        assert _is_malformed_pdf_error(PermissionError("read failed")) is False

    def test_is_anydoc_available_false_in_shipped_state(self):
        """anydoc is not installed in this venv (and not on PyPI — removed
        from requirements.txt). This is the fallback-contract precondition
        the whole tool now rides on."""
        if _anydoc_installed():  # pragma: no cover - only on a dev rig
            pytest.skip("anydoc installed in this env — shipped-state pin not applicable")
        from core.extract.anydoc_extractor import is_anydoc_available

        assert is_anydoc_available() is False

    def test_extract_with_anydoc_missing_file_raises(self, tmp_path):
        from core.extract.anydoc_extractor import extract_with_anydoc

        with pytest.raises(FileNotFoundError, match="Input file not found"):
            extract_with_anydoc(str(tmp_path / "nope.pdf"))

    def test_extract_with_anydoc_empty_path_raises(self):
        from core.extract.anydoc_extractor import extract_with_anydoc

        with pytest.raises(FileNotFoundError, match="Empty file path"):
            extract_with_anydoc("")


# ── Layer 1B: anydoc API dispatch + repair flow (fake anydoc modules) ───────


def _install_fake_anydoc(monkeypatch, to_markdown=None, convert=None, document_cls=None):
    """Inject a fake ``anydoc`` package into ``sys.modules``.

    anydoc is not installable (not on PyPI), so the dispatch logic is pinned
    against hermetic fakes — the shapes mirror the API surface documented at
    core/extract/anydoc_extractor.py:224-236.
    """
    fake = types.ModuleType("anydoc")
    if to_markdown is not None:
        fake.to_markdown = to_markdown
    if convert is not None:
        fake.convert = convert
    if document_cls is not None:
        fake.Document = document_cls
    monkeypatch.setitem(sys.modules, "anydoc", fake)
    return fake


class TestAnydocApiDispatch:
    """``_try_anydoc_extract`` API preference order + result unwrapping."""

    def _extract(self, tmp_path, **kw):

        from core.extract.anydoc_extractor import _try_anydoc_extract

        src = tmp_path / "doc.docx"
        src.write_bytes(b"PK\x03\x04stub")
        return _try_anydoc_extract(str(src), None), src

    def test_prefers_to_markdown_over_convert(self, tmp_path, monkeypatch):
        calls: list[str] = []

        def to_markdown(p):
            calls.append("to_markdown")
            return "MD-A", 3

        def convert(p):  # pragma: no cover - must not be called
            calls.append("convert")
            return "MD-B", 9

        _install_fake_anydoc(monkeypatch, to_markdown=to_markdown, convert=convert)
        (md, pages), _ = self._extract(tmp_path)
        assert (md, pages) == ("MD-A", 3)
        assert calls == ["to_markdown"]

    def test_unwraps_tuple_result(self, tmp_path, monkeypatch):
        _install_fake_anydoc(monkeypatch, to_markdown=lambda p: ("TUP-MD", 5))
        (md, pages), _ = self._extract(tmp_path)
        assert (md, pages) == ("TUP-MD", 5)

    def test_unwraps_dict_result(self, tmp_path, monkeypatch):
        _install_fake_anydoc(monkeypatch, to_markdown=lambda p: {"markdown": "DICT-MD", "pages": 2})
        (md, pages), _ = self._extract(tmp_path)
        assert (md, pages) == ("DICT-MD", 2)

    def test_plain_str_result_zero_pages_then_markers(self, tmp_path, monkeypatch):
        _install_fake_anydoc(
            monkeypatch, to_markdown=lambda p: "x\n<!-- page 1 -->\n<!-- page 2 -->"
        )
        (md, pages), _ = self._extract(tmp_path)
        assert pages == 2  # best-effort from markers

    def test_falls_back_to_convert_api(self, tmp_path, monkeypatch):
        _install_fake_anydoc(monkeypatch, convert=lambda p: ("CONV-MD", 4))
        (md, pages), _ = self._extract(tmp_path)
        assert (md, pages) == ("CONV-MD", 4)

    def test_falls_back_to_document_api(self, tmp_path, monkeypatch):
        class Doc:
            def __init__(self, path):
                self.page_count = 6

            def to_markdown(self):
                return "DOC-MD"

        _install_fake_anydoc(monkeypatch, document_cls=Doc)
        (md, pages), _ = self._extract(tmp_path)
        assert (md, pages) == ("DOC-MD", 6)

    def test_no_usable_api_raises_runtime_error(self, tmp_path, monkeypatch):
        _install_fake_anydoc(monkeypatch, to_markdown=None)
        fake = sys.modules["anydoc"]
        fake.unrelated = 1
        from core.extract.anydoc_extractor import _try_anydoc_extract

        src = tmp_path / "doc.docx"
        src.write_bytes(b"stub")
        with pytest.raises(RuntimeError, match="no usable API"):
            _try_anydoc_extract(str(src), None)

    def test_empty_content_raises_runtime_error(self, tmp_path, monkeypatch):
        _install_fake_anydoc(monkeypatch, to_markdown=lambda p: "   \n  ")
        from core.extract.anydoc_extractor import _try_anydoc_extract

        src = tmp_path / "doc.docx"
        src.write_bytes(b"stub")
        with pytest.raises(RuntimeError, match="empty content"):
            _try_anydoc_extract(str(src), None)


# ── Fixtures ────────────────────────────────────────────────────────────────

_MARKER_TAIL = " body text long enough to clear the fifty character minimum threshold."


def _build_pdf(tmp_path: Path, n_pages: int = 3) -> Path:
    """A real multi-page PDF with per-page marker text (fitz)."""
    import fitz

    p = tmp_path / f"doc{n_pages}.pdf"
    doc = fitz.open()
    for i in range(1, n_pages + 1):
        page = doc.new_page(width=400, height=300)
        page.insert_text((36, 50), f"PDFMARK-{i}{_MARKER_TAIL}")
    doc.save(str(p))
    doc.close()
    return p


def _truncated_pdf(tmp_path: Path, src: Path) -> Path:
    """Cut the xref/trailer tail off a valid PDF (content streams intact).

    The truncation point is content-agnostic: everything from the final
    ``startxref`` keyword onward is dropped, which leaves pypdf unable to
    parse (``PdfStreamError: Stream has ended unexpectedly``) while pikepdf
    can still rebuild the file.
    """
    data = src.read_bytes()
    cut = data.rfind(b"startxref")
    assert cut > 0, "fixture builder: no startxref found in fitz-built PDF"
    out = tmp_path / "truncated.pdf"
    out.write_bytes(data[:cut])
    return out


def _build_docx(tmp_path: Path) -> Path:
    import docx

    p = tmp_path / "notes.docx"
    doc = docx.Document()
    doc.add_heading("Fixture Heading", level=1)
    doc.add_paragraph(f"DOCXMARK{_MARKER_TAIL}")
    doc.save(str(p))
    return p


def _build_epub(tmp_path: Path) -> Path:
    from ebooklib import epub

    p = tmp_path / "book.epub"
    book = epub.EpubBook()
    book.set_title("Fixture Book")
    book.set_language("en")
    ch = epub.EpubHtml(title="Chapter One", file_name="chap1.xhtml", lang="en")
    ch.content = (
        f"<html><body><h1>EPUBMARK chapter one</h1><p>EPUBMARK{_MARKER_TAIL}</p></body></html>"
    )
    book.add_item(ch)
    book.toc = (epub.Link("chap1.xhtml", "Chapter One", "chap1"),)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", ch]
    epub.write_epub(str(p), book)
    return p


def _build_txt(tmp_path: Path) -> Path:
    p = tmp_path / "notes.txt"
    p.write_text(f"TXTMARK{_MARKER_TAIL}", encoding="utf-8")
    return p


def _build_xlsx(tmp_path: Path) -> Path:
    import openpyxl

    p = tmp_path / "sheet.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "XLSXMARK"
    ws["B1"] = "value cell with enough content to clear the minimum length"
    wb.save(str(p))
    return p


def _build_pptx(tmp_path: Path) -> Path:
    from pptx import Presentation

    p = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "PPTXMARK title"
    slide.placeholders[1].text = _MARKER_TAIL.strip()
    prs.save(str(p))
    return p


# ── Layer 1C: engine selection / format validation (dialog, no worker) ─────


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
@pytest.mark.skipif(_anydoc_installed(), reason="anydoc installed — shipped-fallback shape differs")
class TestEngineSelectionCharacterization:
    """Pin the engine-selection contract, including THE key fallback:
    requirements.txt:15-18 says the dialog auto-falls back to markitdown
    when anydoc is absent (the shipped state — anydoc is not on PyPI)."""

    def _make_dialog(self):
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])
        from ui_qt.views.service_dialogs.dialogs.quick_convert_dialog import (
            QuickConvertDialog,
        )

        return QuickConvertDialog(_StubController()), app

    def test_default_engine_falls_back_to_markitdown_when_anydoc_absent(self):
        dlg, _ = self._make_dialog()
        assert dlg.engine_combo.currentIndex() == 1
        assert dlg._selected_engine() == "markitdown"
        assert dlg.engine_combo.itemText(0) == "AnyDoc (not installed)"

    def test_selected_engine_uses_combo_data(self):
        dlg, _ = self._make_dialog()
        dlg.engine_combo.setCurrentIndex(0)
        assert dlg._selected_engine() == "anydoc"
        dlg.engine_combo.setCurrentIndex(1)
        assert dlg._selected_engine() == "markitdown"

    def test_ext_sets_in_sync_with_service_module(self):
        """Drift guard: the dialog duplicates the service sets — a silent
        divergence would split format support between validation and
        conversion. Includes the V0.6.0 (D2) .pptx/.xlsx restoration."""
        from core.extract.anydoc_extractor import (
            ANYDOC_EXTENSIONS,
            MARKITDOWN_EXTENSIONS,
        )
        from ui_qt.views.service_dialogs.dialogs.quick_convert_dialog import (
            QuickConvertDialog,
        )

        assert set(ANYDOC_EXTENSIONS) == QuickConvertDialog._ANYDOC_EXTS
        assert set(MARKITDOWN_EXTENSIONS) == QuickConvertDialog._MARKITDOWN_EXTS

    def test_markitdown_exts_cover_office_formats_after_d2(self):
        """V0.6.0 (D2): markitdown[all] (a core requirements.txt dep) fully
        converts .pptx/.xlsx; with anydoc unshippable these formats
        hard-failed before the fix. Both engines still cover the card's
        advertised PDF/DOCX/EPUB trio."""
        from ui_qt.views.service_dialogs.dialogs.quick_convert_dialog import (
            QuickConvertDialog,
        )

        md = QuickConvertDialog._MARKITDOWN_EXTS
        assert {".pdf", ".docx", ".epub"} <= md
        assert {".pptx", ".xlsx"} <= md  # the D2 fix
        # legacy binary office formats stay anydoc-only (markitdown can't)
        assert ".doc" not in md and ".ppt" not in md and ".xls" not in md

    def test_validate_engine_format_supported_combos(self, tmp_path):
        dlg, _ = self._make_dialog()
        for engine_idx, ext in ((1, ".pdf"), (1, ".docx"), (1, ".epub"), (1, ".pptx")):
            dlg.engine_combo.setCurrentIndex(engine_idx)
            dlg.pdf_edit.setText(str(tmp_path / f"file{ext}"))
            assert dlg._validate_engine_format(silent=False) is True, (engine_idx, ext)
        assert dlg.action_btn.isEnabled()

    def test_validate_engine_format_unsupported_ext_warns_and_disables(self, tmp_path):
        dlg, _ = self._make_dialog()
        dlg.pdf_edit.setText(str(tmp_path / "legacy.doc"))
        assert dlg._validate_engine_format(silent=False) is False
        assert "⚠️ markitdown does not support .doc" in dlg.status_label.text()
        assert not dlg.action_btn.isEnabled()

    def test_validate_engine_format_silent_keeps_button_state(self, tmp_path):
        dlg, _ = self._make_dialog()
        dlg.pdf_edit.setText(str(tmp_path / "legacy.doc"))
        assert dlg._validate_engine_format(silent=True) is False
        assert dlg.action_btn.isEnabled()  # silent: no button touch

    def test_validate_engine_format_no_file_is_true(self):
        dlg, _ = self._make_dialog()
        dlg.pdf_edit.clear()
        assert dlg._validate_engine_format(silent=False) is True

    def test_do_convert_unsupported_ext_hard_blocks_no_worker(self, tmp_path):
        dlg, _ = self._make_dialog()
        src = tmp_path / "weird.xyz"
        src.write_text("x" * 80)
        dlg.pdf_edit.setText(str(src))
        dlg.out_edit.setText(str(tmp_path / "out.md"))
        dlg._do_convert()
        assert "❌ markitdown does not support .xyz files." in dlg.status_label.text()
        assert dlg._worker is None
        assert not (tmp_path / "out.md").exists()

    def test_do_convert_extensionless_file_message_names_this(self, tmp_path):
        """V0.6.0 (D3): a file with NO extension used to render
        ``"does not support  files."`` (empty ext, double space) — now the
        ``{ext or 'this'}` shape from _validate_engine_format applies."""
        dlg, _ = self._make_dialog()
        src = tmp_path / "extensionless"
        src.write_text("x" * 80)
        dlg.pdf_edit.setText(str(src))
        dlg.out_edit.setText(str(tmp_path / "out.md"))
        dlg._do_convert()
        assert "❌ markitdown does not support this files." in dlg.status_label.text()
        assert dlg._worker is None


# ── Layer 2: service-level end-to-end through _convert_work ─────────────────


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
@pytest.mark.skipif(not _MARKITDOWN_REAL, reason="real markitdown package unavailable")
class TestConvertWorkEndToEnd:
    """The real extraction + write on real fixtures — via the
    worker-thread callable (widget-free by design — the W-b5 contract)."""

    def _make_dialog(self):
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])
        from ui_qt.views.service_dialogs.dialogs.quick_convert_dialog import (
            QuickConvertDialog,
        )

        return QuickConvertDialog(_StubController()), app

    def _convert(self, tmp_path, src: Path, out_name="out.md", engine="markitdown", out_ext=".md"):
        dlg, _ = self._make_dialog()
        out = tmp_path / out_name
        result = dlg._convert_work(str(src), str(out), engine, src.suffix.lower(), out_ext)
        return result, out

    # ── the fallback engine (the shipped state: anydoc absent) ──

    def test_pdf_to_markdown_writes_text_and_page_markers(self, tmp_path):
        src = _build_pdf(tmp_path, 3)
        result, out = self._convert(tmp_path, src)
        assert result["success"] is True
        assert "Engine: markitdown" in result["status"]
        assert "Pages: 3" in result["status"]
        text = out.read_text(encoding="utf-8")
        assert "<!-- page 1 -->" in text and "<!-- page 3 -->" in text
        assert "PDFMARK-2" in text
        assert str(out) in result["status"]  # "Saved to:"

    def test_pdf_to_markdown_creates_nested_output_dirs(self, tmp_path):
        src = _build_pdf(tmp_path, 2)
        dlg, _ = self._make_dialog()
        out = tmp_path / "a" / "b" / "out.md"
        result = dlg._convert_work(str(src), str(out), "markitdown", ".pdf", ".md")
        assert result["success"] is True
        assert out.exists()

    def test_docx_to_markdown_recovers_heading_and_body(self, tmp_path):
        result, out = self._convert(tmp_path, _build_docx(tmp_path))
        assert result["success"] is True
        text = out.read_text(encoding="utf-8")
        assert "Fixture Heading" in text
        assert "DOCXMARK" in text

    def test_epub_to_markdown_recovers_chapter_text(self, tmp_path):
        result, out = self._convert(tmp_path, _build_epub(tmp_path))
        assert result["success"] is True
        text = out.read_text(encoding="utf-8")
        assert "EPUBMARK" in text

    def test_txt_to_markdown_through_markitdown(self, tmp_path):
        result, out = self._convert(tmp_path, _build_txt(tmp_path))
        assert result["success"] is True
        assert "TXTMARK" in out.read_text(encoding="utf-8")

    def test_d2_xlsx_to_markdown_through_fallback_engine(self, tmp_path):
        """V0.6.0 (D2): .xlsx converts on the markitdown engine — before the
        fix this hard-failed at validation (anydoc-only list) even though
        markitdown[all] fully supports spreadsheets."""
        result, out = self._convert(tmp_path, _build_xlsx(tmp_path))
        assert result["success"] is True
        assert "XLSXMARK" in out.read_text(encoding="utf-8")

    def test_d2_pptx_to_markdown_through_fallback_engine(self, tmp_path):
        result, out = self._convert(tmp_path, _build_pptx(tmp_path))
        assert result["success"] is True
        assert "PPTXMARK" in out.read_text(encoding="utf-8")

    def test_docx_output_format_rider_writes_real_docx(self, tmp_path):
        """The Format dropdown's .docx rider goes through
        core.outputs.service_formatters.write_output."""
        src = _build_docx(tmp_path)
        result, out = self._convert(tmp_path, src, out_name="out.docx", out_ext=".docx")
        assert result["success"] is True
        import docx as docx_lib

        paras = [p.text for p in docx_lib.Document(str(out)).paragraphs if p.text.strip()]
        assert any("DOCXMARK" in p for p in paras)

    # ── error handling ──

    def test_forced_anydoc_engine_while_absent_errors_cleanly(self, tmp_path):
        src = _build_pdf(tmp_path, 2)
        result, out = self._convert(tmp_path, src, engine="anydoc")
        assert result["success"] is False
        assert "AnyDoc is not installed" in result["status"]
        assert "markitdown" in result["status"]  # actionable advice
        assert not out.exists()

    def test_missing_input_file_errors_nothing_written(self, tmp_path):
        result, out = self._convert(tmp_path, tmp_path / "nope.pdf")
        assert result["success"] is False
        assert result["status"].startswith("❌")
        assert not out.exists()

    def test_tiny_pdf_little_text_guard(self, tmp_path):
        """A near-empty extraction must fail (the scanned-document guard)
        and write nothing."""
        import fitz

        tiny = tmp_path / "tiny.pdf"
        doc = fitz.open()
        page = doc.new_page(width=200, height=200)
        page.insert_text((20, 30), "ab")
        doc.save(str(tiny))
        doc.close()
        result, out = self._convert(tmp_path, tiny)
        assert result["success"] is False
        assert "Very little text extracted" in result["status"]
        assert not out.exists()

    def test_encrypted_pdf_errors_cleanly(self, tmp_path):
        import pypdf

        plain = tmp_path / "plain.pdf"
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=200, height=200)
        with open(plain, "wb") as fh:
            writer.write(fh)
        enc = tmp_path / "enc.pdf"
        reader = pypdf.PdfReader(str(plain))
        w2 = pypdf.PdfWriter()
        w2.append(reader)
        w2.encrypt("fixture-secret")
        with open(enc, "wb") as fh:
            w2.write(fh)
        result, out = self._convert(tmp_path, enc)
        assert result["success"] is False
        assert result["status"].startswith("❌")
        assert not out.exists()

    def test_truncated_pdf_on_fallback_engine_errors_cleanly(self, tmp_path):
        """The markitdown engine has no repair pass (that's the anydoc
        engine's feature) — the malformed PDF surfaces as a clean error."""
        src = _build_pdf(tmp_path, 1)
        trunc = _truncated_pdf(tmp_path, src)
        result, out = self._convert(tmp_path, trunc)
        assert result["success"] is False
        assert result["status"].startswith("❌")
        assert not out.exists()


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
@pytest.mark.skipif(not _fitz_available(), reason="pymupdf unavailable (fixture builder)")
class TestAnydocRepairFlow:
    """The malformed-PDF → pikepdf-repair → retry path, driven with a fake
    anydoc module backed by REAL pypdf (mirrors "anydoc uses pdfminer/pypdf
    under the hood") on a REAL truncated PDF + REAL pikepdf repair."""

    def _install_pypdf_backed_fake(self, monkeypatch):
        import pypdf

        def to_markdown(path):
            reader = pypdf.PdfReader(path)
            parts = []
            for i, page in enumerate(reader.pages):
                parts.append(f"<!-- page {i + 1} -->\n" + (page.extract_text() or ""))
            md = "\n\n".join(parts)
            if not md.strip():
                raise RuntimeError("no text extracted")
            return md, len(reader.pages)

        return _install_fake_anydoc(monkeypatch, to_markdown=to_markdown)

    def test_d1_truncated_pdf_is_repaired_and_text_recovered(self, tmp_path, monkeypatch):
        """V0.6.0 (D1): a truncated PDF (the pypdf ``Stream has ended
        unexpectedly`` shape) previously propagated straight out — the
        pikepdf repair never fired. Now: repair → retry → text recovered."""
        self._install_pypdf_backed_fake(monkeypatch)
        from core.extract.anydoc_extractor import extract_with_anydoc, is_anydoc_available

        assert is_anydoc_available() is True  # the fake is live
        src = _build_pdf(tmp_path, 1)
        trunc = _truncated_pdf(tmp_path, src)
        logs: list[str] = []
        md, pages = extract_with_anydoc(str(trunc), log_callback=logs.append)
        assert "PDFMARK-1" in md
        assert pages == 1
        assert any("pikepdf repair" in line for line in logs)
        assert any("repaired-PDF extraction succeeded" in line for line in logs)

    def test_non_malformed_pdf_error_propagates_without_repair(self, tmp_path, monkeypatch):
        """An encrypted-PDF error (not malformed) must propagate — no
        pointless repair attempt."""
        from pypdf.errors import FileNotDecryptedError

        def to_markdown(path):
            raise FileNotDecryptedError("File has not been decrypted")

        _install_fake_anydoc(monkeypatch, to_markdown=to_markdown)
        from core.extract.anydoc_extractor import extract_with_anydoc

        src = _build_pdf(tmp_path, 1)
        logs: list[str] = []
        with pytest.raises(FileNotDecryptedError):
            extract_with_anydoc(str(src), log_callback=logs.append)
        assert not any("repair" in line for line in logs)


# ── Layer 3: dialog-level drive (the real worker thread + queued signals) ──


def _spin_until_conversion_done(app, dlg, timeout_ms: int = 20000) -> str:
    """Spin the event loop until a final status lands.

    ``_do_convert`` synchronously sets the transient ``Extracting with …``
    status before spawning the worker; every completion path (success /
    error result / worker failure) replaces it. Guard failures (⚠️/❌) land
    synchronously before the first spin.
    """
    import time

    waited = 0
    while waited < timeout_ms:
        app.processEvents()
        status = dlg.status_label.text()
        if status and not status.startswith("Extracting with"):
            return status
        time.sleep(0.02)
        waited += 20
    return dlg.status_label.text()


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
@pytest.mark.skipif(not _fitz_available(), reason="pymupdf unavailable (fixture builder)")
@pytest.mark.skipif(not _MARKITDOWN_REAL, reason="real markitdown package unavailable")
class TestConvertDialogEndToEndDrive:
    """Drive ``_do_convert`` exactly as the UI does: the SimpleWorker
    thread, the queued signals, ``_on_convert_finished`` on the main thread,
    the modal (monkeypatched to a recorder — ``msg.exec()`` at
    dialogs/shared.py:111 would block a headless test forever)."""

    def test_full_pdf_flow_lands_success_status(self, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])

        import ui_qt.views.service_dialogs.dialogs.quick_convert_dialog as qcm

        recorded: list[tuple[str, str, str | None]] = []
        monkeypatch.setattr(
            qcm,
            "_show_success_modal",
            lambda parent, title, message, file_path=None, on_try_another=None: recorded.append(
                (title, message, file_path)
            ),
        )

        dlg = qcm.QuickConvertDialog(_StubController())
        assert dlg._selected_engine() == "markitdown"  # the shipped fallback

        src = _build_pdf(tmp_path, 3)
        out = tmp_path / "driven.md"
        dlg.pdf_edit.setText(str(src))
        dlg.fmt_combo.setCurrentIndex(1)  # explicit ".md"
        dlg.out_edit.setText(str(out))
        dlg._do_convert()

        status = _spin_until_conversion_done(app, dlg)
        assert status.startswith("✅"), status
        assert f"Saved to: {out}" in status
        assert "Pages: 3" in status
        assert "Engine: markitdown" in status

        text = out.read_text(encoding="utf-8")
        assert "PDFMARK-2" in text and "<!-- page 1 -->" in text

        # The modal fired with the real output path (recorded, not exec'd).
        assert recorded and recorded[0][0] == "📄 Document → Markdown Complete"
        assert recorded[0][2] == str(out)
        assert dlg.action_btn.text() == "Close"
        # Headless: isVisible() needs a shown parent; isHidden() reflects the
        # explicit setVisible() flip in _on_convert_finished.
        assert not dlg.again_btn.isHidden()

    def test_full_docx_flow_lands_success_status(self, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])

        import ui_qt.views.service_dialogs.dialogs.quick_convert_dialog as qcm

        monkeypatch.setattr(qcm, "_show_success_modal", lambda *a, **k: None)
        dlg = qcm.QuickConvertDialog(_StubController())
        src = _build_docx(tmp_path)
        out = tmp_path / "driven_docx.md"
        dlg.pdf_edit.setText(str(src))
        dlg.fmt_combo.setCurrentIndex(1)
        dlg.out_edit.setText(str(out))
        dlg._do_convert()

        status = _spin_until_conversion_done(app, dlg)
        assert status.startswith("✅"), status
        assert "Fixture Heading" in out.read_text(encoding="utf-8")

    def test_missing_paths_guard_no_worker_spawned(self, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QApplication

        QApplication.instance() or QApplication([])

        import ui_qt.views.service_dialogs.dialogs.quick_convert_dialog as qcm

        monkeypatch.setattr(qcm, "_show_success_modal", lambda *a, **k: None)
        dlg = qcm.QuickConvertDialog(_StubController())
        dlg._do_convert()  # no input, no output
        assert dlg.status_label.text() == "⚠️ Please pick a document file first."
        assert dlg._worker is None

        dlg.pdf_edit.setText(str(_build_pdf(tmp_path, 2)))
        dlg._do_convert()  # input but no output
        assert dlg.status_label.text() == "⚠️ Please choose where to save the file."
        assert dlg._worker is None

    def test_unsupported_ext_flow_lands_error_no_worker(self, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QApplication

        QApplication.instance() or QApplication([])

        import ui_qt.views.service_dialogs.dialogs.quick_convert_dialog as qcm

        dlg = qcm.QuickConvertDialog(_StubController())
        src = tmp_path / "weird.xyz"
        src.write_text("x" * 80)
        dlg.pdf_edit.setText(str(src))
        dlg.out_edit.setText(str(tmp_path / "out.md"))
        dlg._do_convert()
        assert "❌ markitdown does not support .xyz files." in dlg.status_label.text()
        assert dlg._worker is None
        assert not (tmp_path / "out.md").exists()

    def test_forced_anydoc_engine_flow_lands_error_and_reenables(self, tmp_path, monkeypatch):
        """A user who re-selects the (uninstalled) AnyDoc engine gets a
        clean worker-run error, not a crash — and can retry."""
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])

        import ui_qt.views.service_dialogs.dialogs.quick_convert_dialog as qcm

        monkeypatch.setattr(qcm, "_show_success_modal", lambda *a, **k: None)
        dlg = qcm.QuickConvertDialog(_StubController())
        dlg.engine_combo.setCurrentIndex(0)  # force the anydoc engine
        assert dlg._selected_engine() == "anydoc"

        src = _build_pdf(tmp_path, 2)
        out = tmp_path / "forced.md"
        dlg.pdf_edit.setText(str(src))
        dlg.fmt_combo.setCurrentIndex(1)
        dlg.out_edit.setText(str(out))
        dlg._do_convert()

        status = _spin_until_conversion_done(app, dlg)
        assert "AnyDoc is not installed" in status
        assert dlg.action_btn.text() == "Convert"
        assert dlg.action_btn.isEnabled()
        assert not out.exists()


# ── Stub controller (dialogs only need get/update_setting — verified
#    via grep: the only controller.* call sites are shared.py) ───────────


class _StubController:
    def __init__(self) -> None:
        self._settings: dict[str, object] = {}

    def get_setting(self, key: str, default=None):
        return self._settings.get(key, default)

    def update_setting(self, key: str, value) -> None:
        self._settings[key] = value
