"""Phase 2 (Test-Net Deepening) — long-tail coverage for keypoints_builder.py.

Baseline: 264 stmts / 96 missed (64%). The existing suite builds keypoints
only through the full converter (happy path). This file covers the missing
branches directly:

  * dedup: rapidfuzz-unavailable fallback, embedder-exception fallback,
    the embedding-dedup path (sim > 0.85), short-bullet skip, the
    >15 dropped-details log cap, "no unique keypoints" early return
  * topic ordering fallback (order_bullets_by_topic crash)
  * segments (source text + page markers) and judge_results (Layer-8
    evidence lookup → 🛡️ verified badge + confidence bump)
  * metadata-computation exception → the plain fallback metadata
  * disk-full (errno 28) re-raise with the actionable message
  * ``_build_docx`` (legacy) + ``_build_docx_enhanced`` /
    ``_add_bullet_to_docx`` (confidence tiers, verified badge, excerpts)

Hermetic: python-docx + rapidfuzz + numpy are all installed in the test
venv; LLM/embedder boundaries are monkeypatched fakes (no network). All
file writes go to pytest's tmp_path.

Run: ``COVERAGE_FILE=.coverage.p2c3 python3 -m pytest tests/test_p2_sumtail_keypoints.py -q --no-cov``
"""

from __future__ import annotations

import errno
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

_project_root = Path(__file__).resolve().parent.parent
if str(_project_root) not in sys.path:
    sys.path.insert(0, str(_project_root))

from core.models import SummaryResult
from core.summarisation.keypoints_builder import KeypointsBuilder


class _Log:
    def __init__(self):
        self.calls = []

    def __call__(self, msg, level="INFO"):
        self.calls.append((msg, level))

    def infos(self):
        return [m for m, lvl in self.calls if lvl == "INFO"]

    def debugs(self):
        return [m for m, lvl in self.calls if lvl == "DEBUG"]


@pytest.fixture(autouse=True)
def _no_embedder(monkeypatch):
    """Default: no embedding model → the build takes the rapidfuzz dedup
    branch (hermetic — the real singleton would probe localhost:11434).
    The embedding-dedup test overrides this with its own monkeypatch."""
    import core.llm.model_singletons as ms

    monkeypatch.setattr(ms, "get_embed_model", lambda: None)


def _summary(bullets, heading="Chapter One", success=True):
    return SummaryResult(
        chapter_heading=heading,
        executive_summary="An executive summary.",
        bullet_points=list(bullets),
        thorough_summary="A thorough summary.",
        success=success,
    )


# ── 1. Early returns ──────────────────────────────────────────────────────


class TestEarlyReturns:
    def test_no_successful_summaries_returns_none(self, tmp_path):
        kb = KeypointsBuilder()
        failed = _summary(["b"], success=False)
        assert kb.build([failed], "Book", tmp_path) == (None, None)
        assert kb.final_keypoint_count == 0

    def test_no_bullets_returns_none(self, tmp_path):
        kb = KeypointsBuilder()
        assert kb.build([_summary([])], "Book", tmp_path) == (None, None)

    def test_all_bullets_too_short_returns_none(self, tmp_path):
        """Bullets shorter than 2 chars are filtered → no unique keypoints."""
        kb = KeypointsBuilder()
        log = _Log()
        kb.log = log
        out = kb.build([_summary(["x", "y"])], "Book", tmp_path)
        assert out == (None, None)
        assert any("No unique keypoints" in i for i in log.infos())


# ── 2. Dedup paths ────────────────────────────────────────────────────────


class TestDedupPaths:
    def test_exact_duplicates_removed(self, tmp_path):
        kb = KeypointsBuilder()
        log = _Log()
        kb.log = log
        dup = "Orexin neurons regulate wakefulness across species."
        md, docx = kb.build([_summary([dup, dup, dup])], "Book", tmp_path)
        assert md is not None and docx is not None
        assert kb.final_keypoint_count == 1
        assert any("removed 2 duplicate(s)" in i for i in log.infos())

    def test_rapidfuzz_near_duplicates_removed(self, tmp_path):
        """token_set_ratio ≥ 75 near-duplicates are merged (V3.9.2 75% rule)."""
        kb = KeypointsBuilder()
        md, _ = kb.build(
            [
                _summary(
                    [
                        "The hypothalamus regulates hunger through orexin neurons.",
                        "The hypothalamus regulates hunger via orexin neurons.",
                    ]
                )
            ],
            "Book",
            tmp_path,
        )
        content = md.read_text(encoding="utf-8")
        # Near-duplicate merged → exactly ONE keypoint mentions the hypothalamus
        assert content.count("hypothalamus") == 1
        assert kb.final_keypoint_count == 1

    def test_dropped_details_capped_at_15(self, tmp_path):
        """>15 dropped duplicates → 15 detail lines + 'and N more'."""
        kb = KeypointsBuilder()
        log = _Log()
        kb.log = log
        # 20 lexically DISTINCT bullets (low token_set_ratio between pairs so
        # only the exact duplicates below are dropped)
        words = [
            "aardvarks forage at dawn",
            "quokkas hop through scrub",
            "narwhals dive beneath ice",
            "ibexes climb sheer cliffs",
            "pangolins curl into armor",
            "okapis browse jungle shadows",
            "axolotls regrow their limbs",
            "sifakas sideways through trees",
            "kakapos boom across valleys",
            "tuataras bask on offshore rocks",
            "vaquitas surface at dusk",
            "saigas roam the open steppe",
            "geladas graze high meadows",
            "fennecs dig deep burrows",
            "tapirs wade along riverbanks",
            "manatees drift through springs",
            "lynxes stalk northern forests",
            "otters juggle smooth pebbles",
            "ravens solve complex puzzles",
            "octopuses unscrew jar lids",
        ]
        bullets = [f"{w[0].upper()}{w[1:]}." for w in words]
        # 17 exact duplicates of the FIRST bullet
        bullets = bullets + [bullets[0]] * 17
        kb.build([_summary(bullets)], "Book", tmp_path)
        details = [m for m in log.debugs() if "(exact match)" in m]
        assert len(details) == 15
        assert any("and 2 more" in m for m in log.debugs())

    def test_rapidfuzz_unavailable_exact_dedup_still_works(self, tmp_path, monkeypatch):
        """rapidfuzz import blocked → _has_rf False → only exact dedup runs.

        The two bullets normalize differently (so exact dedup keeps both) but
        rapidfuzz token_set_ratio would merge them (≥ 75) — with rapidfuzz
        blocked BOTH survive.
        """
        monkeypatch.setitem(sys.modules, "rapidfuzz", None)
        kb = KeypointsBuilder()
        md, _ = kb.build(
            [
                _summary(
                    [
                        "Orexin neurons regulate mammalian wakefulness cycles.",
                        "Orexin neurons regulate mammalian wakefulness patterns.",
                    ]
                )
            ],
            "Book",
            tmp_path,
        )
        # Without rapidfuzz the near-duplicate survives
        assert kb.final_keypoint_count == 2

    def test_embed_model_exception_falls_back_to_rapidfuzz(self, tmp_path, monkeypatch):
        """A crashing get_embed_model → logger.debug → rapidfuzz dedup path."""
        import core.llm.model_singletons as ms

        def _boom():
            raise RuntimeError("embed backend down")

        monkeypatch.setattr(ms, "get_embed_model", _boom)
        kb = KeypointsBuilder()
        dup = "The exact same bullet text for the embedding fallback test."
        md, _ = kb.build([_summary([dup, dup])], "Book", tmp_path)
        assert kb.final_keypoint_count == 1

    def test_embedding_dedup_merges_semantic_duplicates(self, tmp_path, monkeypatch):
        """The V3.3.0 C3 embedding path: cosine > 0.85 → semantic duplicate."""
        import numpy as np

        import core.llm.embedding_cache as ec
        import core.llm.model_singletons as ms

        fake_model = SimpleNamespace(name="fake-embedder")

        # Two orthogonal unit vectors: vec A for bullets 0+1 (identical →
        # sim 1.0 → dup), vec B for bullet 2 (orthogonal → kept)
        vec_a = np.array([1.0, 0.0])
        vec_b = np.array([0.0, 1.0])

        def fake_cached_encode(model, texts):
            embs = []
            for t in texts:
                # Semantic-duplicate bullets map to vec_a regardless of text
                embs.append(vec_a if "rapport" in t or "greeting" in t else vec_b)
            return np.array(embs)

        monkeypatch.setattr(ms, "get_embed_model", lambda: fake_model)
        monkeypatch.setattr(ec, "cached_encode", fake_cached_encode)

        kb = KeypointsBuilder()
        log = _Log()
        kb.log = log
        bullets = [
            "Use signature greetings to create rapport with clients.",
            "A signature greeting technique builds immediate rapport.",
            "Orexin neurons control the sleep wake cycle.",
        ]
        md, _ = kb.build([_summary(bullets)], "Book", tmp_path)
        assert kb.final_keypoint_count == 2
        assert any("embedding sim > 0.85" in m for m in log.debugs())


# ── 3. Topic ordering + metadata fallbacks ────────────────────────────────


class TestOrderingAndMetadata:
    def test_topic_ordering_crash_falls_back_to_input_order(self, tmp_path, monkeypatch):
        import core.summarisation.text_preprocessor as tp

        def _boom(bullets):
            raise RuntimeError("spacy model missing")

        monkeypatch.setattr(tp, "order_bullets_by_topic", _boom)
        kb = KeypointsBuilder()
        bullets = ["First keypoint about identity.", "Second keypoint about memory."]
        md, _ = kb.build([_summary(bullets)], "Book", tmp_path)
        content = md.read_text(encoding="utf-8")
        # Order preserved (fallback), both present
        assert content.index("First keypoint") < content.index("Second keypoint")

    def test_metadata_crash_falls_back_to_plain_metadata(self, tmp_path, monkeypatch):
        """compute_bullet_confidence raising → fallback dicts (confidence 50,
        classification 'additional'). Segments are supplied so the confidence
        call is actually reached (it's gated behind ``if source_text``)."""
        import core.summarisation.text_preprocessor as tp

        def _boom(bp, src, concepts):
            raise RuntimeError("confidence model unavailable")

        monkeypatch.setattr(tp, "compute_bullet_confidence", _boom)
        kb = KeypointsBuilder()
        log = _Log()
        kb.log = log
        seg = SimpleNamespace(content="source text", start_page=1)
        md, docx = kb.build(
            [_summary(["A perfectly ordinary keypoint."])], "Book", tmp_path, segments=[seg]
        )
        assert md is not None and docx is not None

    def test_classification_counts_logged(self, tmp_path, monkeypatch):
        """The K1+Classification INFO line reports core/supporting/additional."""
        import core.summarisation.text_preprocessor as tp

        labels = {"b-core": "core", "b-supp": "supporting", "b-add": "additional"}

        def _classify(bp, concepts, source_text):
            return labels.get(bp.strip(), "additional")

        monkeypatch.setattr(tp, "classify_bullet", _classify)
        kb = KeypointsBuilder()
        log = _Log()
        kb.log = log
        kb.build([_summary(["b-core", "b-supp", "b-add"])], "Book", tmp_path)
        line = next(m for m in log.infos() if "K1+Classification" in m)
        assert "1 core, 1 supporting, 1 additional" in line


# ── 4. Segments + judge_results (Layer 8 evidence) ────────────────────────


class TestSegmentsAndEvidence:
    def test_segments_provide_source_text_and_pages(self, tmp_path, monkeypatch):
        import core.summarisation.text_preprocessor as tp

        seen_pages = []

        def _find_page(bp, source_text, page_markers):
            seen_pages.append(page_markers)
            return 7

        monkeypatch.setattr(tp, "find_source_page", _find_page)
        kb = KeypointsBuilder()
        seg1 = SimpleNamespace(content="First segment about identity.", start_page=3)
        seg2 = SimpleNamespace(content="Second segment about memory.", start_page=7)
        bullet = "Identity forms through identification with others."
        md, docx = kb.build([_summary([bullet])], "Book", tmp_path, segments=[seg1, seg2])
        assert md is not None and docx is not None
        assert seen_pages and seen_pages[0][0] == (3, 3)

    def test_judge_evidence_upgrades_to_verified(self, tmp_path, monkeypatch):
        """A JudgeResult with bullet_evidence → 🛡️ verified + confidence ≥ 85."""
        import core.summarisation.text_preprocessor as tp

        monkeypatch.setattr(tp, "compute_bullet_confidence", lambda bp, src, c: 42)
        kb = KeypointsBuilder()
        log = _Log()
        kb.log = log
        bullet = "Identity forms through identification with others."
        jr = SimpleNamespace(
            grounding_layer8_available=True,
            bullet_evidence={bullet: "the infant identifies with the imago"},
        )
        md, docx = kb.build(
            [_summary([bullet])],
            "Book",
            tmp_path,
            segments=[
                SimpleNamespace(content="the infant identifies with the imago", start_page=1)
            ],
            judge_results=[jr],
        )
        assert md is not None and docx is not None
        line = next(m for m in log.infos() if "K1+Classification" in m)
        assert "1 Layer-8 verified" in line
        # The verified badge + source citation are in the DOCX
        from docx import Document

        doc = Document(str(docx))
        all_text = "\n".join(p.text for p in doc.paragraphs)
        assert "🛡️ Verified source" in all_text
        assert "the infant identifies with the imago" in all_text

    def test_judge_none_entries_skipped(self, tmp_path):
        """None entries in judge_results are skipped without crash."""
        kb = KeypointsBuilder()
        md, docx = kb.build(
            [_summary(["A plain keypoint without evidence."])],
            "Book",
            tmp_path,
            judge_results=[None, None],
        )
        assert md is not None and docx is not None

    def test_l8_inactive_no_verified_suffix(self, tmp_path):
        """judge_results without grounding_layer8_available → no 🛡️ suffix."""
        kb = KeypointsBuilder()
        bullet = "A plain keypoint without evidence."
        jr = SimpleNamespace(grounding_layer8_available=False, bullet_evidence={})
        log = _Log()
        kb.log = log
        kb.build([_summary([bullet])], "Book", tmp_path, judge_results=[jr])
        line = next(m for m in log.infos() if "K1+Classification" in m)
        assert "Layer-8 verified" not in line

    def test_excerpt_and_page_rendered_in_docx(self, tmp_path, monkeypatch):
        """K2/H2: a source excerpt + page citation render as the 📖 sub-bullet."""
        import core.summarisation.text_preprocessor as tp

        monkeypatch.setattr(
            tp, "find_source_excerpt", lambda bp, src: "the verbatim source excerpt"
        )
        monkeypatch.setattr(tp, "find_source_page", lambda bp, src, pm: 12)
        kb = KeypointsBuilder()
        seg = SimpleNamespace(content="the verbatim source excerpt lives here", start_page=12)
        md, docx = kb.build(
            [_summary(["A keypoint with an excerpt."])], "Book", tmp_path, segments=[seg]
        )
        from docx import Document

        doc = Document(str(docx))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "📖 Source:" in text
        assert '"the verbatim source excerpt"' in text

    def test_excerpt_without_page_uses_plain_source_label(self, tmp_path, monkeypatch):
        """An excerpt but NO page → still the 📖 label (V10.24.0 removed the
        unreliable page citation from the label)."""
        import core.summarisation.text_preprocessor as tp

        monkeypatch.setattr(tp, "find_source_excerpt", lambda bp, src: "excerpt without page")
        monkeypatch.setattr(tp, "find_source_page", lambda bp, src, pm: None)
        kb = KeypointsBuilder()
        seg = SimpleNamespace(content="excerpt without page appears here", start_page=1)
        md, docx = kb.build(
            [_summary(["A keypoint without a page."])], "Book", tmp_path, segments=[seg]
        )
        from docx import Document

        doc = Document(str(docx))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "📖 Source:" in text
        assert '"excerpt without page"' in text

    def test_page_marker_construction_crash_swallowed(self, tmp_path):
        """A segment whose start_page property raises (non-AttributeError) →
        the page-marker try/except logs and continues (no page citations)."""

        class _BadSegment:
            content = "Segment content for the bad page marker."

            @property
            def start_page(self):
                raise ValueError("corrupt page marker")

        kb = KeypointsBuilder()
        md, docx = kb.build(
            [_summary(["A keypoint for the bad segment."])],
            "Book",
            tmp_path,
            segments=[_BadSegment()],
        )
        assert md is not None and docx is not None


# ── 5. Disk-full handling (V10.4.0 Phase 6A-4) ───────────────────────────


class TestDiskFull:
    def test_errno_28_raises_actionable_disk_full(self, tmp_path, monkeypatch):
        import utils.atomic_io as aio

        def _raise_no_space(path, content, **kw):
            raise OSError(errno.ENOSPC, "No space left on device")

        monkeypatch.setattr(aio, "atomic_write_text", _raise_no_space)
        kb = KeypointsBuilder()
        with pytest.raises(OSError, match="Disk is full"):
            kb.build([_summary(["A keypoint that cannot be written."])], "Book", tmp_path)

    def test_other_oserror_reraised_unchanged(self, tmp_path, monkeypatch):
        import utils.atomic_io as aio

        def _raise_perm(path, content, **kw):
            raise OSError(errno.EACCES, "Permission denied")

        monkeypatch.setattr(aio, "atomic_write_text", _raise_perm)
        kb = KeypointsBuilder()
        with pytest.raises(OSError, match="Permission denied"):
            kb.build([_summary(["A keypoint that cannot be written."])], "Book", tmp_path)


# ── 6. Output formats ─────────────────────────────────────────────────────


class TestOutputFormats:
    def test_markdown_structure(self, tmp_path):
        kb = KeypointsBuilder()
        bullets = ["First keypoint.", "Second keypoint."]
        md, _ = kb.build([_summary(bullets)], "My Book", tmp_path)
        content = md.read_text(encoding="utf-8")
        assert content.startswith("# Keypoints: My Book")
        assert "**Total keypoints:** 2" in content
        assert "1. First keypoint." in content
        assert "2. Second keypoint." in content

    def test_legacy_build_docx(self, tmp_path):
        """The legacy _build_docx produces a valid .docx with List Bullets."""
        from docx import Document

        kb = KeypointsBuilder()
        path = tmp_path / "legacy.docx"
        kb._build_docx("Legacy Book", ["Legacy bullet one."], path)
        doc = Document(str(path))
        assert any("Legacy bullet one." in p.text for p in doc.paragraphs)

    def test_enhanced_docx_confidence_tiers(self, tmp_path, monkeypatch):
        """All three confidence indicators (🟢🟡🔴) render in the DOCX.

        Segments are supplied so ``source_text`` is non-empty — the
        confidence call is gated behind ``if source_text else 50``.
        """
        import core.summarisation.text_preprocessor as tp

        conf = {"high": 90, "mid": 60, "low": 20}
        monkeypatch.setattr(
            tp,
            "compute_bullet_confidence",
            lambda bp, src, c: conf.get(bp.strip(), 50),
        )
        kb = KeypointsBuilder()
        seg = SimpleNamespace(content="source segment text", start_page=1)
        md, docx = kb.build([_summary(["high", "mid", "low"])], "Book", tmp_path, segments=[seg])
        from docx import Document

        doc = Document(str(docx))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "🟢" in text
        assert "🟡" in text
        assert "🔴" in text
        assert "[90/100]" in text
        assert "[20/100]" in text

    def test_output_directory_created_lazily(self, tmp_path):
        """V9.3.1: the output dir is created even when nobody made it first."""
        kb = KeypointsBuilder()
        nested = tmp_path / "deeply" / "nested" / "dir"
        md, docx = kb.build([_summary(["A keypoint in a nested dir."])], "Book", nested)
        assert md.exists()
        assert docx.exists()

    def test_final_keypoint_count_is_post_dedup(self, tmp_path):
        """V18.20.3.10 (W11): the count reflects what SHIPPED, not raw."""
        kb = KeypointsBuilder()
        dup = "Deterministic duplicate bullet for the count test."
        kb.build([_summary([dup, dup, "Unique bullet."])], "Book", tmp_path)
        assert kb.final_keypoint_count == 2
