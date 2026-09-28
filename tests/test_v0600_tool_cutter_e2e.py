"""V0.6.0 (the dedicated-tools verification campaign) — Tool 1: the
PDF/DOCX Cutter, end-to-end.

Inventory + trace (evidence-first):
- Card: ``pdf_docx_cutter`` — ``ui_qt/views/tool_dashboard.py:41`` TOOL_CARDS
  row 3 (EXTRACT).
- Routing: ``ui_qt/views/convert_view.py:208`` ``_on_tool_clicked`` →
  ``PdfDocxCutterDialog``.
- Dialog: ``ui_qt/views/service_dialogs/dialogs/pdf_docx_cutter_dialog.py``
  (the V18.20.4.6 O-b2 split; service layer is EMBEDDED — ``_cut_pdf``
  pypdf at :274, ``_cut_docx`` python-docx at :294, ``_parse_range`` :159).
- Existing coverage: source-pin only (test_v182045_p4a_service_dialog_threads.py
  / test_v182046_p5_dialog_split.py assert the worker wiring + result-dict
  keys via AST) — ZERO functional tests. This file is the first net.

Layers (house UI-test pattern, cf. test_v182038_audit_fixes):
1. Pure characterization of ``_parse_range`` (no Qt).
2. Service-level end-to-end: real fitz-built PDFs and python-docx-built
   DOCX fixtures cut through the worker-thread callable ``_cut_work``
   (widget-free by design — the W-b5 contract).
3. Dialog-level drive: the real SimpleWorker thread + queued signals +
   ``_on_cut_finished`` on the main thread (offscreen QApplication, the
   modal monkeypatched to a recorder so ``msg.exec()`` cannot block).
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

import pytest

_project_root = Path(__file__).resolve().parent.parent
if str(_project_root) not in sys.path:
    sys.path.insert(0, str(_project_root))


# ── Layer 0: availability gates ─────────────────────────────────────────────


def _qt_available() -> bool:
    try:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication  # noqa: F401

        return True
    except Exception:
        return False


def _fitz_available() -> bool:
    try:
        import fitz  # noqa: F401

        return True
    except Exception:
        return False


def _docx_available() -> bool:
    try:
        import docx  # noqa: F401

        return True
    except Exception:
        return False


# ── Layer 1: _parse_range characterization (pure, no Qt) ────────────────────


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
class TestParseRangeCharacterization:
    """Pin the exact current semantics of the range parser.

    C1 context (V0.6.0): a NON-EMPTY unparseable range (``"abc"``,
    ``"5-"``, ``"0"``) still returns None at the PARSER level, but since
    V0.6.0 ``_do_cut`` refuses to run on that shape (status warning, no
    worker) — the silent whole-document cut is closed at the dialog layer.
    Typographic dashes are normalized BEFORE parsing (the pasted-from-PDF
    "2–4" family).
    """

    def _dialog_cls(self):
        from ui_qt.views.service_dialogs.dialogs.pdf_docx_cutter_dialog import (
            PdfDocxCutterDialog,
        )

        return PdfDocxCutterDialog

    def test_all_and_empty_mean_none(self):
        assert self._dialog_cls()._parse_range("all") is None
        assert self._dialog_cls()._parse_range("") is None
        assert self._dialog_cls()._parse_range("  ALL  ") is None

    def test_simple_ranges(self):
        assert self._dialog_cls()._parse_range("1-3") == [1, 2, 3]
        assert self._dialog_cls()._parse_range("2-2") == [2]
        assert self._dialog_cls()._parse_range("1,3,5") == [1, 3, 5]
        assert self._dialog_cls()._parse_range("1, 3, 5") == [1, 3, 5]

    def test_reversed_range_is_normalized(self):
        assert self._dialog_cls()._parse_range("3-1") == [1, 2, 3]

    def test_zero_clamped_to_one(self):
        assert self._dialog_cls()._parse_range("0-2") == [1, 2]

    def test_duplicates_collapse_sorted(self):
        assert self._dialog_cls()._parse_range("5,3,5,3") == [3, 5]

    def test_unparseable_parts_skipped(self):
        assert self._dialog_cls()._parse_range("1,abc,3") == [1, 3]

    # C1 pins (parser-level None shapes — the _do_cut guard is the fix)
    def test_c1_garbage_means_none_parser_level(self):
        assert self._dialog_cls()._parse_range("abc") is None

    def test_c1_open_ended_range_means_none_parser_level(self):
        assert self._dialog_cls()._parse_range("5-") is None

    def test_c1_zero_page_means_none_parser_level(self):
        assert self._dialog_cls()._parse_range("0") is None

    # V0.6.0 (C1 fix): typographic dash normalization — ranges pasted from
    # PDFs/TOCs arrive with en/em/minus dashes.
    def test_endash_range_parses(self):
        assert self._dialog_cls()._parse_range("2\u20134") == [2, 3, 4]

    def test_emdash_range_parses(self):
        assert self._dialog_cls()._parse_range("2\u20144") == [2, 3, 4]

    def test_minus_sign_range_parses(self):
        assert self._dialog_cls()._parse_range("2\u22124") == [2, 3, 4]

    def test_endash_list_mixed_with_ascii(self):
        assert self._dialog_cls()._parse_range("1\u20133, 5, 7\u20138") == [1, 2, 3, 5, 7, 8]


# ── Fixtures ────────────────────────────────────────────────────────────────


def _build_pdf(tmp_path: Path, n_pages: int) -> Path:
    """A real multi-page PDF with per-page marker text (fitz)."""
    import fitz

    p = tmp_path / f"doc{n_pages}.pdf"
    doc = fitz.open()
    for i in range(1, n_pages + 1):
        page = doc.new_page(width=400, height=300)
        page.insert_text((36, 50), f"PAGE-{i}")
    doc.save(str(p))
    doc.close()
    return p


def _build_docx_pagebreaks(tmp_path: Path, n_pages: int) -> Path:
    """A real DOCX with EXPLICIT page breaks (strategy 1 — w:br type=page)."""
    import docx
    from docx.enum.text import WD_BREAK

    p = tmp_path / f"pbdoc{n_pages}.docx"
    doc = docx.Document()
    for i in range(1, n_pages + 1):
        para = doc.add_paragraph(f"DPAGE-{i}")
        para.add_run(f" tail-{i}")
        if i < n_pages:
            para.runs[0].add_break(WD_BREAK.PAGE)
    doc.save(str(p))
    return p


def _build_docx_many_paras(tmp_path: Path, n_paras: int) -> Path:
    """A real DOCX with NO breaks (strategy 3 — ~40 paragraphs per page)."""
    import docx

    p = tmp_path / f"mpdoc{n_paras}.docx"
    doc = docx.Document()
    for i in range(n_paras):
        doc.add_paragraph(f"PARA-{i}")
    doc.save(str(p))
    return p


# ── Layer 2: service-level end-to-end through _cut_work ─────────────────────


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
@pytest.mark.skipif(not _fitz_available(), reason="pymupdf unavailable (fixture builder)")
class TestCutPdfEndToEnd:
    """The real pypdf cut on real PDFs — via the worker-thread callable."""

    def _make_dialog(self):
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])
        from ui_qt.views.service_dialogs.dialogs.pdf_docx_cutter_dialog import (
            PdfDocxCutterDialog,
        )

        return PdfDocxCutterDialog(_StubController()), app

    def _cut(self, tmp_path, pages_text: str, n_pages: int = 10):
        dlg, _ = self._make_dialog()
        src = _build_pdf(tmp_path, n_pages)
        out = tmp_path / "out.pdf"
        result = dlg._cut_work(str(src), str(out), dlg._parse_range(pages_text))
        return result, out

    def test_range_cut_selects_exact_pages(self, tmp_path):
        result, out = self._cut(tmp_path, "2-4")
        assert result == {"out_path": str(out)}
        import pypdf

        r = pypdf.PdfReader(str(out))
        assert len(r.pages) == 3
        texts = [p.extract_text() for p in r.pages]
        assert all("PAGE-2" in texts[0] for _ in [0])
        assert "PAGE-2" in texts[0] and "PAGE-3" in texts[1] and "PAGE-4" in texts[2]
        assert "PAGE-1" not in texts[0]

    def test_list_cut_selects_disjoint_pages(self, tmp_path):
        result, out = self._cut(tmp_path, "1,3,5")
        import pypdf

        r = pypdf.PdfReader(str(out))
        assert len(r.pages) == 3
        texts = [p.extract_text() for p in r.pages]
        assert "PAGE-1" in texts[0] and "PAGE-3" in texts[1] and "PAGE-5" in texts[2]

    def test_all_cut_copies_every_page(self, tmp_path):
        result, out = self._cut(tmp_path, "all")
        import pypdf

        r = pypdf.PdfReader(str(out))
        assert len(r.pages) == 10

    def test_out_of_range_only_selection_errors_and_writes_nothing(self, tmp_path):
        """C2 (V0.6.0): a selection entirely beyond the page count used to
        write a 0-page PDF and report success — now it errors through the
        standard worker result and nothing is written."""
        result, out = self._cut(tmp_path, "99")
        assert "error" in result and "no pages matched the range" in result["error"]
        assert "❌ Error:" in result["error"]
        assert not out.exists()

    def test_mixed_in_and_out_of_range_keeps_valid_pages(self, tmp_path):
        result, out = self._cut(tmp_path, "9,99")
        import pypdf

        r = pypdf.PdfReader(str(out))
        assert len(r.pages) == 1
        assert "PAGE-9" in r.pages[0].extract_text()

    def test_unsupported_extension_returns_error_result(self, tmp_path):
        dlg, _ = self._make_dialog()
        src = tmp_path / "doc.txt"
        src.write_text("not a pdf")
        result = dlg._cut_work(str(src), str(tmp_path / "out.txt"), None)
        assert result["error"] == "❌ Only .pdf and .docx are supported."

    def test_missing_input_file_returns_error_result(self, tmp_path):
        dlg, _ = self._make_dialog()
        result = dlg._cut_work(str(tmp_path / "nope.pdf"), str(tmp_path / "out.pdf"), None)
        assert "error" in result and "❌ Error:" in result["error"]

    def test_output_parent_dirs_created(self, tmp_path):
        dlg, _ = self._make_dialog()
        src = _build_pdf(tmp_path, 3)
        out = tmp_path / "a" / "b" / "out.pdf"
        result = dlg._cut_work(str(src), str(out), None)
        assert result == {"out_path": str(out)}
        assert out.exists()


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
@pytest.mark.skipif(not _docx_available(), reason="python-docx unavailable")
class TestCutDocxEndToEnd:
    """The real python-docx cut on real DOCX fixtures — via _cut_work."""

    def _make_dialog(self):
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])
        from ui_qt.views.service_dialogs.dialogs.pdf_docx_cutter_dialog import (
            PdfDocxCutterDialog,
        )

        return PdfDocxCutterDialog(_StubController()), app

    def _read_paras(self, path: Path) -> list[str]:
        import docx

        return [p.text for p in docx.Document(str(path)).paragraphs if p.text.strip()]

    def test_pagebreak_doc_cut_page_2(self, tmp_path):
        """Strategy 1 (explicit page breaks): 3 'pages', keep the middle."""
        dlg, _ = self._make_dialog()
        src = _build_docx_pagebreaks(tmp_path, 3)
        out = tmp_path / "out.docx"
        result = dlg._cut_work(str(src), str(out), [2])
        assert result == {"out_path": str(out)}
        paras = self._read_paras(out)
        assert paras == ["DPAGE-2 tail-2"]

    def test_pagebreak_doc_cut_all(self, tmp_path):
        dlg, _ = self._make_dialog()
        src = _build_docx_pagebreaks(tmp_path, 3)
        out = tmp_path / "out.docx"
        assert dlg._cut_work(str(src), str(out), None) == {"out_path": str(out)}
        paras = self._read_paras(out)
        assert [p for p in paras if p.startswith("DPAGE-")] == [
            "DPAGE-1 tail-1",
            "DPAGE-2 tail-2",
            "DPAGE-3 tail-3",
        ]

    def test_paragraph_heuristic_doc_cut_page_1(self, tmp_path):
        """Strategy 3 (no breaks, ~40 paras/page): 80 paras → 2 pages."""
        dlg, _ = self._make_dialog()
        src = _build_docx_many_paras(tmp_path, 80)
        out = tmp_path / "out.docx"
        result = dlg._cut_work(str(src), str(out), [1])
        assert result == {"out_path": str(out)}
        paras = self._read_paras(out)
        assert len(paras) == 40
        assert paras[0] == "PARA-0" and paras[-1] == "PARA-39"

    def test_ranged_cut_keeps_run_formatting(self, tmp_path):
        """Bold/italic/underline/font survive the copy (V11.31.1 contract)."""
        import docx

        src = tmp_path / "fmt.docx"
        doc = docx.Document()
        para = doc.add_paragraph()
        run = para.add_run("BOLDITALIC")
        run.bold = True
        run.italic = True
        para2 = doc.add_paragraph()
        run2 = para2.add_run("tail")
        run2.add_break(docx.enum.text.WD_BREAK.PAGE)
        doc.save(str(src))
        dlg, _ = self._make_dialog()
        out = tmp_path / "out.docx"
        assert dlg._cut_work(str(src), str(out), [1]) == {"out_path": str(out)}
        paras = docx.Document(str(out)).paragraphs
        assert paras[0].runs[0].bold is True
        assert paras[0].runs[0].italic is True

    def test_tables_copied_on_full_cut_dropped_on_ranged_cut(self, tmp_path):
        """V11.31.1 documented contract: tables ride along only on 'all'."""
        import docx

        for pages in (None, [1]):
            src = tmp_path / f"tbl_{'all' if pages is None else 'rng'}.docx"
            doc = docx.Document()
            doc.add_paragraph("P1")
            doc.add_paragraph("P2")
            t = doc.add_table(rows=1, cols=2)
            t.rows[0].cells[0].text = "CELL-A"
            t.rows[0].cells[1].text = "CELL-B"
            doc.save(str(src))
            dlg, _ = self._make_dialog()
            out = tmp_path / f"out_{'all' if pages is None else 'rng'}.docx"
            result = dlg._cut_work(str(src), str(out), pages)
            assert result == {"out_path": str(out)}
            new_doc = docx.Document(str(out))
            if pages is None:
                assert len(new_doc.tables) == 1
                assert new_doc.tables[0].rows[0].cells[0].text == "CELL-A"
            else:
                assert len(new_doc.tables) == 0


# ── Layer 3: dialog-level drive (the real worker thread + queued signals) ──


@pytest.mark.skipif(not _qt_available(), reason="PySide6/offscreen unavailable")
@pytest.mark.skipif(not _fitz_available(), reason="pymupdf unavailable (fixture builder)")
class TestCutDialogEndToEndDrive:
    """Drive ``_do_cut`` exactly as the UI does: the SimpleWorker thread,
    the queued signals, ``_on_cut_finished`` on the main thread, the modal
    (monkeypatched to a recorder — ``msg.exec()`` must never block a test)."""

    def test_full_cut_flow_lands_success_status(self, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])

        import ui_qt.views.service_dialogs.dialogs.pdf_docx_cutter_dialog as cutter_mod

        recorded: list[tuple[str, str, str | None]] = []
        monkeypatch.setattr(
            cutter_mod,
            "_show_success_modal",
            lambda parent, title, message, file_path=None, on_try_another=None: recorded.append(
                (title, message, file_path)
            ),
        )

        dlg = cutter_mod.PdfDocxCutterDialog(_StubController())
        src = _build_pdf(tmp_path, 6)
        out = tmp_path / "driven.pdf"
        dlg.in_edit.setText(str(src))
        dlg.out_edit.setText(str(out))
        dlg.range_edit.setText("2-3")

        dlg._do_cut()

        # Spin the event loop until the worker finishes + queued slots run.
        deadline_ms = 15000
        waited = 0
        import time

        while waited < deadline_ms:
            app.processEvents()
            if dlg.status_label.text().startswith(("✅", "❌", "⚠️")):
                break
            time.sleep(0.02)
            waited += 20

        assert dlg.status_label.text().startswith("✅"), dlg.status_label.text()
        assert f"Saved to: {out}" in dlg.status_label.text()

        import pypdf

        r = pypdf.PdfReader(str(out))
        assert len(r.pages) == 2
        assert "PAGE-2" in r.pages[0].extract_text()

        # The modal fired with the real output path (recorded, not exec'd).
        assert recorded and recorded[0][0] == "✂ Cut Complete"
        assert recorded[0][2] == str(out)
        # The button flips to Close on success (V18.20.4.5 W-b5 contract).
        assert dlg.action_btn.text() == "Close"

    def test_missing_paths_guard_no_worker_spawned(self, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QApplication

        QApplication.instance() or QApplication([])

        import ui_qt.views.service_dialogs.dialogs.pdf_docx_cutter_dialog as cutter_mod

        dlg = cutter_mod.PdfDocxCutterDialog(_StubController())
        dlg._do_cut()  # no input, no output
        assert dlg.status_label.text() == "⚠️ Please pick an input file."
        assert dlg._worker is None

        dlg.in_edit.setText(str(_build_pdf(tmp_path, 2)))
        dlg._do_cut()  # input but no output
        assert dlg.status_label.text() == "⚠️ Please choose where to save the output."
        assert dlg._worker is None

    def test_c1_garbage_range_guards_no_worker_spawned(self, tmp_path, monkeypatch):
        """V0.6.0 (C1): an unparseable NON-EMPTY range warns and refuses to
        run — the silent whole-document cut is closed."""
        from PySide6.QtWidgets import QApplication

        QApplication.instance() or QApplication([])

        import ui_qt.views.service_dialogs.dialogs.pdf_docx_cutter_dialog as cutter_mod

        dlg = cutter_mod.PdfDocxCutterDialog(_StubController())
        dlg.in_edit.setText(str(_build_pdf(tmp_path, 4)))
        dlg.out_edit.setText(str(tmp_path / "guard.pdf"))
        dlg.range_edit.setText("abc")
        dlg._do_cut()
        assert dlg.status_label.text().startswith("⚠️ Couldn't read that page range")
        assert dlg._worker is None
        assert not (tmp_path / "guard.pdf").exists()

        dlg.range_edit.setText("0")
        dlg._do_cut()
        assert dlg.status_label.text().startswith("⚠️ Couldn't read that page range")
        assert dlg._worker is None

    def test_c1_open_ended_range_guards_no_worker_spawned(self, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QApplication

        QApplication.instance() or QApplication([])

        import ui_qt.views.service_dialogs.dialogs.pdf_docx_cutter_dialog as cutter_mod

        dlg = cutter_mod.PdfDocxCutterDialog(_StubController())
        dlg.in_edit.setText(str(_build_pdf(tmp_path, 4)))
        dlg.out_edit.setText(str(tmp_path / "guard.pdf"))
        dlg.range_edit.setText("5-")
        dlg._do_cut()
        assert dlg.status_label.text().startswith("⚠️ Couldn't read that page range")
        assert dlg._worker is None

    def test_endash_range_cuts_through_the_worker(self, tmp_path, monkeypatch):
        """V0.6.0 (C1): the pasted-from-PDF en-dash range now cuts the
        requested pages (was: silent whole-document cut)."""
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])

        import ui_qt.views.service_dialogs.dialogs.pdf_docx_cutter_dialog as cutter_mod

        monkeypatch.setattr(cutter_mod, "_show_success_modal", lambda *a, **k: None)
        dlg = cutter_mod.PdfDocxCutterDialog(_StubController())
        src = _build_pdf(tmp_path, 6)
        out = tmp_path / "endash.pdf"
        dlg.in_edit.setText(str(src))
        dlg.out_edit.setText(str(out))
        dlg.range_edit.setText("2\u20134")
        dlg._do_cut()

        import time

        waited = 0
        while waited < 15000:
            app.processEvents()
            if dlg.status_label.text().startswith(("✅", "❌", "⚠️")):
                break
            time.sleep(0.02)
            waited += 20

        assert dlg.status_label.text().startswith("✅"), dlg.status_label.text()
        import pypdf

        r = pypdf.PdfReader(str(out))
        assert len(r.pages) == 3
        assert "PAGE-2" in r.pages[0].extract_text()
        assert "PAGE-4" in r.pages[2].extract_text()

    def test_unsupported_ext_flow_lands_error_status_and_reenables(self, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QApplication

        app = QApplication.instance() or QApplication([])

        import ui_qt.views.service_dialogs.dialogs.pdf_docx_cutter_dialog as cutter_mod

        recorded: list[tuple] = []
        monkeypatch.setattr(cutter_mod, "_show_success_modal", lambda *a, **k: recorded.append(a))

        dlg = cutter_mod.PdfDocxCutterDialog(_StubController())
        src = tmp_path / "doc.txt"
        src.write_text("x")
        out = tmp_path / "out.txt"
        dlg.in_edit.setText(str(src))
        dlg.out_edit.setText(str(out))
        dlg.range_edit.setText("all")
        dlg._do_cut()

        import time

        waited = 0
        while waited < 15000:
            app.processEvents()
            if dlg.status_label.text().startswith(("✅", "❌", "⚠️")):
                break
            time.sleep(0.02)
            waited += 20

        assert "Only .pdf and .docx are supported." in dlg.status_label.text()
        assert dlg.action_btn.text() == "Cut" and dlg.action_btn.isEnabled()
        assert not recorded


# ── Stub controller (dialogs only need get/update_setting — verified
#    via grep: the only controller.* call sites are shared.py:643,646) ───────


class _StubController:
    def __init__(self) -> None:
        self._settings: dict[str, object] = {}

    def get_setting(self, key: str, default=None):
        return self._settings.get(key, default)

    def update_setting(self, key: str, value) -> None:
        self._settings[key] = value
